from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import logging
import re
import textwrap
from urllib.parse import urlencode

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.mail import send_mail
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q, Sum, Value
from django.db.models.functions import Coalesce, Lower, Trim, TruncMonth, TruncWeek
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect
from django.template.loader import render_to_string
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.http import url_has_allowed_host_and_scheme
from django.views import View
from django.views.generic import CreateView, ListView, RedirectView, TemplateView, UpdateView
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .access_control import (
    ModulePermissionRequiredMixin,
    TRACKER_PERMISSION_ITEMS,
    build_permission_items,
    get_first_accessible_route_name,
    get_user_permission_record,
)
from .expense_categories import (
    COUNTER_EXPENSE_CATEGORY,
    OFFICE_EXPENSE_CATEGORY,
    ensure_expense_categories_for_role,
    get_expense_category_purpose_map,
    get_expense_category_options as get_saved_expense_category_options,
    get_or_create_role_expense_category,
    get_or_create_role_expense_purpose,
    normalize_expense_category_name,
    normalize_expense_purpose_name,
)
from .income_categories import (
    DEFAULT_INCOME_CATEGORY,
    INCOME_CATEGORY_COUNTER,
    INCOME_CATEGORY_OFFICE,
    normalize_income_category_name,
)
from .income_purposes import (
    get_income_category_purpose_map,
    get_or_create_role_income_purpose,
    normalize_income_purpose_name,
)
from .forms import (
    DailyCashSettlementForm,
    ExpenseCategoryForm,
    ExpenseForm,
    ExpensePurposeForm,
    IncomeForm,
    IncomePurposeForm,
    PurchaseForm,
    PurchasePaymentForm,
    ReconciliationOpeningBalanceForm,
    SupplierForm,
    UserManagementForm,
)
from .models import (
    DailyCashSettlement,
    ExpenseCategory,
    ExpensePurpose,
    ExpenseRecord,
    IncomeRecord,
    IncomePurpose,
    PaymentMethod,
    PurchasePayment,
    PurchaseRecord,
    ReconciliationIncomeEntry,
    ReconciliationOpeningBalance,
    SalesLedgerRecord,
    SalesPaymentMode,
    Supplier,
    SupplierStatus,
)
from .purchase_helpers import sync_purchase_to_expense
from .purchase_sync import (
    PurchaseSyncError,
    SQLSERVER_PURCHASE_SOURCE_PREFIX,
    sync_purchases_from_sqlserver,
)
from .sales_sync import SalesSyncError, sync_sales_from_sqlserver
from .supplier_sync import SupplierSyncError, sync_suppliers_from_sqlserver
from .user_sync import UserSyncError, sync_users_from_sqlserver
from .user_roles import filter_queryset_by_role, get_user_role_label

User = get_user_model()
RECONCILIATION_NON_CASH_METHODS = (
    PaymentMethod.CARD,
    PaymentMethod.UPI,
    PaymentMethod.BANK_TRANSFER,
    PaymentMethod.OTHER,
)
SETTLEMENT_NON_CASH_EXPENSE_METHODS = tuple(
    method for method in RECONCILIATION_NON_CASH_METHODS if method != PaymentMethod.CARD
)
RECONCILIATION_NON_CASH_METHOD_CASEFOLD = tuple(
    method.casefold() for method in RECONCILIATION_NON_CASH_METHODS
)
RECONCILIATION_HISTORY_PER_PAGE = 20
MANUAL_PURCHASE_SOURCE_PREFIX = "MANUAL:"
logger = logging.getLogger(__name__)
SPLIT_PAYMENT_MODE_FILTER = "Split"

def _sum_field(queryset, field_name):
    return queryset.aggregate(total=Sum(field_name))["total"] or Decimal("0.00")


def _sum_amount(queryset):
    return _sum_field(queryset, "amount")


def _sum_settlement_income(queryset):
    return queryset.aggregate(total=Sum("actual_sales"))["total"] or Decimal("0.00")


def get_sales_cash_from_settlement(settlement):
    return settlement.actual_sales - settlement.gpay_settled


def get_cash_in_hand_amount(
    opening_balance,
    sales_ledger_cash,
    expense_amount,
    counter_income_amount=Decimal("0.00"),
):
    return opening_balance + sales_ledger_cash + counter_income_amount - expense_amount


def get_settlement_closing_balance(
    opening_balance,
    sales_ledger_cash,
    expense_amount,
    cash_settled,
    cash_denomination_total=Decimal("0.00"),
    counter_income_amount=Decimal("0.00"),
):
    cash_in_hand = get_cash_in_hand_amount(
        opening_balance,
        sales_ledger_cash,
        expense_amount,
        counter_income_amount=counter_income_amount,
    )
    cash_difference = get_cash_difference_amount(
        cash_denomination_total,
        cash_in_hand,
    )
    return (cash_in_hand - cash_settled) + cash_difference


def _apply_payment_method_filter(queryset, payment_methods=None):
    if payment_methods is None:
        return queryset
    if isinstance(payment_methods, str):
        return queryset.filter(payment_method=payment_methods)
    return queryset.filter(payment_method__in=payment_methods)


def build_income_ledger_entries_for_period(user, start_date=None, end_date=None):
    entries = []
    income_queryset = filter_queryset_by_role(IncomeRecord.objects.all(), user)
    settlement_queryset = filter_queryset_by_role(
        DailyCashSettlement.objects.all(),
        user,
    )

    if start_date is not None:
        income_queryset = income_queryset.filter(transaction_date__gte=start_date)
        settlement_queryset = settlement_queryset.filter(settlement_date__gte=start_date)
    if end_date is not None:
        income_queryset = income_queryset.filter(transaction_date__lt=end_date)
        settlement_queryset = settlement_queryset.filter(settlement_date__lt=end_date)

    for record in income_queryset.order_by("-transaction_date", "-created_at", "-pk"):
        entries.append(
            {
                "pk": record.pk,
                "transaction_date": record.transaction_date,
                "title": record.title,
                "source": record.source,
                "category": record.category,
                "payment_method": record.payment_method,
                "amount": record.amount,
                "entry_type": "manual",
                "sort_date": record.transaction_date,
                "sort_timestamp": record.created_at,
            }
        )

    for settlement in settlement_queryset.order_by("-settlement_date", "-updated_at", "-pk"):
        settlement_sales_cash = get_sales_cash_from_settlement(settlement)
        entries.append(
            {
                "pk": "",
                "transaction_date": settlement.settlement_date,
                "title": "Daily Settlement Income",
                "source": (
                    f"Cash Rs. {settlement_sales_cash} | "
                    f"GPay Rs. {settlement.gpay_settled}"
                ),
                "category": INCOME_LEDGER_SETTLEMENT_CATEGORY,
                "payment_method": INCOME_LEDGER_SETTLEMENT_PAYMENT_METHOD,
                "amount": settlement.actual_sales,
                "settlement_cash_amount": settlement_sales_cash,
                "settlement_upi_amount": settlement.gpay_settled,
                "entry_type": "settlement",
                "sort_date": settlement.settlement_date,
                "sort_timestamp": settlement.updated_at,
            }
        )

    entries.sort(
        key=lambda item: (item["sort_date"], item["sort_timestamp"]),
        reverse=True,
    )
    return entries


def build_income_ledger_entries(user):
    return build_income_ledger_entries_for_period(user)


INCOME_LEDGER_SETTLEMENT_CATEGORY = "Daily Settlement"
INCOME_LEDGER_SETTLEMENT_PAYMENT_METHOD = "Cash + UPI"


def _clone_income_settlement_entry(entry, payment_method, amount, source):
    filtered_entry = dict(entry)
    filtered_entry["payment_method"] = payment_method
    filtered_entry["amount"] = amount
    filtered_entry["source"] = source
    return filtered_entry


def get_income_entry_for_payment_filter(entry, payment_method=""):
    entry_payment_method = (entry.get("payment_method") or "").strip()

    if not payment_method or payment_method == "All":
        return entry

    if entry.get("entry_type") != "settlement":
        if entry_payment_method == payment_method:
            return entry
        return None

    settlement_cash_amount = entry.get("settlement_cash_amount", Decimal("0.00"))
    settlement_upi_amount = entry.get("settlement_upi_amount", Decimal("0.00"))

    if payment_method == INCOME_LEDGER_SETTLEMENT_PAYMENT_METHOD:
        return entry
    if payment_method == PaymentMethod.CASH and settlement_cash_amount > 0:
        return _clone_income_settlement_entry(
            entry,
            PaymentMethod.CASH,
            settlement_cash_amount,
            f"Cash Rs. {settlement_cash_amount}",
        )
    if payment_method == PaymentMethod.UPI and settlement_upi_amount > 0:
        return _clone_income_settlement_entry(
            entry,
            PaymentMethod.UPI,
            settlement_upi_amount,
            f"GPay Rs. {settlement_upi_amount}",
        )
    return None


def get_raw_income_filter_values(request):
    return {
        "start_date": (request.GET.get("start_date") or "").strip(),
        "end_date": (request.GET.get("end_date") or "").strip(),
        "category": (request.GET.get("category") or "").strip(),
        "payment_method": (request.GET.get("payment_method") or "").strip(),
    }


def get_income_filter_values(request):
    filter_values = get_raw_income_filter_values(request)
    if not any(filter_values.values()):
        today_value = date.today().isoformat()
        filter_values["start_date"] = today_value
        filter_values["end_date"] = today_value
        return filter_values

    if filter_values["start_date"] and not filter_values["end_date"]:
        filter_values["end_date"] = filter_values["start_date"]
    if filter_values["end_date"] and not filter_values["start_date"]:
        filter_values["start_date"] = filter_values["end_date"]
    return filter_values


def apply_income_filters(entries, filter_values):
    start_date = parse_date(filter_values["start_date"])
    end_date = parse_date(filter_values["end_date"])
    category = filter_values["category"]
    payment_method = filter_values["payment_method"]

    if start_date and end_date and start_date > end_date:
        start_date, end_date = end_date, start_date

    filtered_entries = []
    for entry in entries:
        entry_date = entry.get("transaction_date")
        entry_category = (entry.get("category") or "").strip()

        if start_date and entry_date and entry_date < start_date:
            continue
        if end_date and entry_date and entry_date > end_date:
            continue
        if category and category != "All" and entry_category.casefold() != category.casefold():
            continue
        filtered_entry = get_income_entry_for_payment_filter(entry, payment_method)
        if filtered_entry is None:
            continue
        filtered_entries.append(filtered_entry)
    return filtered_entries


def get_income_category_options(entries, selected_category=""):
    default_options = [
        INCOME_CATEGORY_COUNTER,
        INCOME_CATEGORY_OFFICE,
        INCOME_LEDGER_SETTLEMENT_CATEGORY,
    ]
    discovered_options = {
        (entry.get("category") or "").strip()
        for entry in entries
        if (entry.get("category") or "").strip()
    }
    ordered_options = [
        option for option in default_options if option in discovered_options or option != INCOME_LEDGER_SETTLEMENT_CATEGORY
    ]
    ordered_options.extend(
        sorted(
            discovered_options - set(ordered_options),
            key=lambda value: value.lower(),
        )
    )
    if selected_category and selected_category not in ordered_options:
        ordered_options.append(selected_category)
    return ordered_options


def get_income_payment_method_options(entries, selected_payment_method=""):
    default_options = [value for value, _label in PaymentMethod.choices]
    default_options.append(INCOME_LEDGER_SETTLEMENT_PAYMENT_METHOD)
    discovered_options = {
        (entry.get("payment_method") or "").strip()
        for entry in entries
        if (entry.get("payment_method") or "").strip()
    }
    ordered_options = [option for option in default_options if option in discovered_options or option in default_options]
    ordered_options.extend(
        sorted(
            discovered_options - set(ordered_options),
            key=lambda value: value.lower(),
        )
    )
    if selected_payment_method and selected_payment_method not in ordered_options:
        ordered_options.append(selected_payment_method)
    return [(option, option) for option in ordered_options]


def summarize_income_entries(entries):
    manual_total = sum(
        (
            entry.get("amount", Decimal("0.00"))
            for entry in entries
            if entry.get("entry_type") == "manual"
        ),
        Decimal("0.00"),
    )
    settlement_total = sum(
        (
            entry.get("amount", Decimal("0.00"))
            for entry in entries
            if entry.get("entry_type") == "settlement"
        ),
        Decimal("0.00"),
    )
    return {
        "manual_total": manual_total,
        "settlement_total": settlement_total,
        "total": manual_total + settlement_total,
    }


def _shift_month(month_start, months_back):
    year = month_start.year
    month = month_start.month - months_back
    while month <= 0:
        year -= 1
        month += 12
    return date(year, month, 1)


def _next_month(month_start):
    if month_start.month == 12:
        return date(month_start.year + 1, 1, 1)
    return date(month_start.year, month_start.month + 1, 1)


def build_monthly_overview(user, months=6, anchor_date=None):
    anchor_date = anchor_date or date.today()
    month_anchor = anchor_date.replace(day=1)
    rows = []

    for months_back in range(months - 1, -1, -1):
        month_start = _shift_month(month_anchor, months_back)
        month_end = _next_month(month_start)
        if month_start.year == anchor_date.year and month_start.month == anchor_date.month:
            month_end = anchor_date + timedelta(days=1)

        income_total = _sum_amount(
            filter_queryset_by_role(IncomeRecord.objects.all(), user).filter(
                transaction_date__gte=month_start,
                transaction_date__lt=month_end,
            )
        )
        expense_total = _sum_amount(
            filter_queryset_by_role(ExpenseRecord.objects.all(), user).filter(
                transaction_date__gte=month_start,
                transaction_date__lt=month_end,
            )
        )
        rows.append(
            {
                "label": month_start.strftime("%b %Y"),
                "income": income_total,
                "expense": expense_total,
                "balance": income_total - expense_total,
            }
        )

    highest_total = max(
        [max(row["income"], row["expense"]) for row in rows],
        default=Decimal("1.00"),
    )
    if highest_total == 0:
        highest_total = Decimal("1.00")

    for row in rows:
        row["income_width"] = round((row["income"] / highest_total) * 100, 2)
        row["expense_width"] = round((row["expense"] / highest_total) * 100, 2)

    return rows


def get_reporting_income_total(income_queryset, settlement_queryset):
    return _sum_amount(income_queryset) + _sum_settlement_income(settlement_queryset)


def build_reporting_monthly_overview(user, months=6, anchor_date=None):
    anchor_date = anchor_date or date.today()
    month_anchor = anchor_date.replace(day=1)
    window_start = _shift_month(month_anchor, months - 1)
    window_end = min(_next_month(month_anchor), anchor_date + timedelta(days=1))
    income_queryset = filter_queryset_by_role(IncomeRecord.objects.all(), user)
    expense_queryset = filter_queryset_by_role(ExpenseRecord.objects.all(), user)
    settlement_queryset = filter_queryset_by_role(
        DailyCashSettlement.objects.all(),
        user,
    )
    income_totals = build_monthly_total_map(
        income_queryset,
        date_field="transaction_date",
        amount_field="amount",
        start_date=window_start,
        end_date=window_end,
    )
    expense_totals = build_monthly_total_map(
        expense_queryset,
        date_field="transaction_date",
        amount_field="amount",
        start_date=window_start,
        end_date=window_end,
    )
    settlement_totals = build_monthly_total_map(
        settlement_queryset,
        date_field="settlement_date",
        amount_field="actual_sales",
        start_date=window_start,
        end_date=window_end,
    )
    rows = []

    for months_back in range(months - 1, -1, -1):
        month_start = _shift_month(month_anchor, months_back)
        income_total = income_totals.get(month_start, Decimal("0.00")) + settlement_totals.get(
            month_start,
            Decimal("0.00"),
        )
        expense_total = expense_totals.get(month_start, Decimal("0.00"))
        balance = income_total - expense_total
        rows.append(
            {
                "label": month_start.strftime("%b %Y"),
                "income": income_total,
                "expense": expense_total,
                "balance": balance,
                "margin": round((balance / income_total) * 100, 1)
                if income_total
                else Decimal("0.0"),
            }
        )

    highest_total = max(
        [max(row["income"], row["expense"]) for row in rows],
        default=Decimal("1.00"),
    )
    if highest_total == 0:
        highest_total = Decimal("1.00")

    for row in rows:
        row["income_width"] = round((row["income"] / highest_total) * 100, 2)
        row["expense_width"] = round((row["expense"] / highest_total) * 100, 2)

    return rows


def build_reporting_income_category_overview(income_queryset, settlement_queryset):
    category_totals = {}

    for row in income_queryset.values("category").annotate(total=Sum("amount")):
        category = row["category"] or "Uncategorized"
        category_totals[category] = category_totals.get(category, Decimal("0.00")) + (
            row["total"] or Decimal("0.00")
        )

    settlement_total = _sum_settlement_income(settlement_queryset)
    if settlement_total:
        category_totals["Daily Settlement"] = category_totals.get(
            "Daily Settlement",
            Decimal("0.00"),
        ) + settlement_total

    ranked_rows = [
        {"category": category, "total": total}
        for category, total in sorted(
            category_totals.items(),
            key=lambda item: (-item[1], item[0]),
        )[:5]
    ]
    return build_ranked_category_overview(ranked_rows)


def get_report_month(request, parameter_name="report_month"):
    raw_value = (request.GET.get(parameter_name) or "").strip()
    if re.fullmatch(r"\d{4}-\d{2}", raw_value):
        try:
            return date.fromisoformat(f"{raw_value}-01")
        except ValueError:
            pass
    return date.today().replace(day=1)


def get_report_month_bounds(report_month):
    month_start = report_month.replace(day=1)
    return month_start, _next_month(month_start)


def get_report_month_details(report_month):
    month_start, month_end = get_report_month_bounds(report_month)
    today = date.today()
    is_current_month = month_start.year == today.year and month_start.month == today.month
    if is_current_month:
        return {
            "month_start": month_start,
            "month_end": month_end,
            "period_end": today + timedelta(days=1),
            "anchor_date": today,
            "is_partial": True,
            "scope_label": f"Month to date through {today:%B} {today.day}, {today.year}",
        }
    return {
        "month_start": month_start,
        "month_end": month_end,
        "period_end": month_end,
        "anchor_date": month_end - timedelta(days=1),
        "is_partial": False,
        "scope_label": f"Full month of {report_month:%B %Y}",
    }


def build_monthly_total_map(queryset, date_field, amount_field, start_date, end_date):
    filters = {
        f"{date_field}__gte": start_date,
        f"{date_field}__lt": end_date,
    }
    monthly_totals = {}
    for row in (
        queryset.filter(**filters)
        .annotate(month=TruncMonth(date_field))
        .values("month")
        .annotate(total=Sum(amount_field))
        .order_by("month")
    ):
        month_value = row["month"]
        if hasattr(month_value, "date"):
            month_value = month_value.date()
        monthly_totals[month_value] = row["total"] or Decimal("0.00")
    return monthly_totals


def autosize_report_worksheet(worksheet):
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            cell_value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(cell_value))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 32)


def style_report_export_sheet(worksheet, header_row=1):
    header_fill = PatternFill(fill_type="solid", fgColor="163B6D")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font

    worksheet.freeze_panes = worksheet[f"A{header_row + 1}"]
    worksheet.auto_filter.ref = worksheet.dimensions
    autosize_report_worksheet(worksheet)


def build_reports_excel_response(user, params):
    workspace = build_sales_report_workspace(user, params, include_export_records=True)
    filter_values = workspace["sales_report_filters"]
    filtered_summary = workspace["filtered_summary"]
    payment_split_lookup = workspace["payment_split_lookup"]
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet["A1"] = "Mahilmart Sales Report"
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet.append([])
    summary_sheet.append(["Metric", "Value"])
    for label, value in (
        ("Range", filter_values["range_label"]),
        ("Start Date", filter_values["start_date"].isoformat()),
        ("End Date", filter_values["end_date"].isoformat()),
        ("Payment Mode", filter_values["payment_mode"] or "All"),
        ("Scope", filter_values["scope_label"]),
        ("Bill Count", filtered_summary["count"]),
        ("Total Sales", float(filtered_summary["total_amount"])),
        ("Received Amount", float(filtered_summary["received_amount"])),
        ("Pending Credit", float(filtered_summary["balance_amount"])),
        ("Average Bill", float(workspace["average_bill_amount"])),
        ("Cash Collection", float(payment_split_lookup.get("cash", Decimal("0.00")))),
        ("UPI / Card Collection", float(payment_split_lookup.get("upi", Decimal("0.00")))),
        ("Credit Outstanding", float(payment_split_lookup.get("credit", Decimal("0.00")))),
        ("Split Bills", filtered_summary["split_bill_count"]),
    ):
        summary_sheet.append([label, value])
    style_report_export_sheet(summary_sheet, header_row=3)

    for sheet_name, chart in (
        ("Daily Sales", workspace["sales_charts"][0]),
        ("Weekly Sales", workspace["sales_charts"][1]),
        ("Monthly Sales", workspace["sales_charts"][2]),
    ):
        chart_sheet = workbook.create_sheet(sheet_name)
        chart_sheet.append(["Label", "Short Label", "Bill Count", "Sales Total"])
        for row in chart["rows"]:
            chart_sheet.append(
                [
                    row["label"],
                    row["short_label"],
                    row["count"],
                    float(row["total"]),
                ]
            )
        style_report_export_sheet(chart_sheet)

    payment_split_sheet = workbook.create_sheet("Payment Split")
    payment_split_sheet.append(["Payment Type", "Amount", "Share %"])
    for row in workspace["payment_split_rows"]:
        payment_split_sheet.append(
            [
                row["label"],
                float(row["amount"]),
                float(row["share"]),
            ]
        )
    style_report_export_sheet(payment_split_sheet)

    ledger_sheet = workbook.create_sheet("Sales Ledger")
    ledger_sheet.append(
        [
            "Sale Date",
            "Bill No",
            "Customer Name",
            "Payment Mode",
            "Net Amount",
            "Received Amount",
            "Balance Amount",
            "Split Cash",
            "Split Card",
        ]
    )
    for record in workspace["export_records"]:
        ledger_sheet.append(
            [
                record.sale_date.isoformat() if record.sale_date else "",
                record.bill_no,
                record.customer_name or "Walk-in Customer",
                record.display_payment_mode,
                float(record.net_amount),
                float(record.effective_received_amount),
                float(record.effective_balance_amount),
                float(record.split_cash_amount),
                float(record.split_card_amount),
            ]
        )
    style_report_export_sheet(ledger_sheet)

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        "attachment; filename="
        f'"mahilmart_sales_report_{filter_values["start_date"]:%Y%m%d}_{filter_values["end_date"]:%Y%m%d}.xlsx"'
    )
    return response


def get_reconciliation_filter_values(params):
    today = date.today()
    start_date = parse_date((params.get("start_date") or "").strip())
    end_date = parse_date((params.get("end_date") or "").strip())

    if start_date is None and end_date is None:
        start_date = today
        end_date = today
    elif start_date is None:
        start_date = end_date
    elif end_date is None:
        end_date = start_date

    if start_date and end_date and start_date > end_date:
        start_date, end_date = end_date, start_date

    return {
        "start_date": start_date or today,
        "end_date": end_date or today,
    }


def get_reconciliation_history_view(params):
    selected_view = (params.get("view") or "").strip().lower()
    return "expense" if selected_view == "expense" else "income"


def build_reconciliation_url(
    filter_values,
    history_view="income",
    route_name="reconciliation",
):
    resolved_history_view = (
        "expense" if str(history_view).strip().lower() == "expense" else "income"
    )
    return (
        f"{reverse(route_name)}?"
        f"{urlencode(
            {
                'start_date': filter_values['start_date'].isoformat(),
                'end_date': filter_values['end_date'].isoformat(),
                'view': resolved_history_view,
            }
        )}"
    )


def build_reconciliation_redirect_url(params, route_name="reconciliation"):
    filter_values = get_reconciliation_filter_values(params)
    return build_reconciliation_url(
        filter_values,
        get_reconciliation_history_view(params),
        route_name=route_name,
    )


def build_reconciliation_pagination(
    request,
    total_count,
    *,
    page_param="page",
    per_page=RECONCILIATION_HISTORY_PER_PAGE,
    history_view=None,
):
    paginator = Paginator(range(total_count), per_page)
    page_obj = paginator.get_page((request.GET.get(page_param) or "").strip() or 1)
    url_overrides = {}
    if history_view:
        url_overrides["view"] = history_view

    previous_page_url = ""
    next_page_url = ""
    if page_obj.has_previous():
        previous_page_url = build_page_url(
            request,
            page_obj.previous_page_number(),
            page_param=page_param,
            **url_overrides,
        )
    if page_obj.has_next():
        next_page_url = build_page_url(
            request,
            page_obj.next_page_number(),
            page_param=page_param,
            **url_overrides,
        )

    return {
        "page_obj": page_obj,
        "previous_page_url": previous_page_url,
        "next_page_url": next_page_url,
        "pagination_links": [
            {
                "number": page_number,
                "url": build_page_url(
                    request,
                    page_number,
                    page_param=page_param,
                    **url_overrides,
                ),
                "is_current": page_number == page_obj.number,
            }
            for page_number in page_obj.paginator.page_range
        ],
        "showing_from": page_obj.start_index() if total_count else 0,
        "showing_to": page_obj.end_index() if total_count else 0,
    }


def get_reconciliation_income_querysets(user, filter_values):
    manual_queryset = filter_queryset_by_role(
        IncomeRecord.objects.all(),
        user,
    ).filter(
        transaction_date__gte=filter_values["start_date"],
        transaction_date__lte=filter_values["end_date"],
    )
    settlement_queryset = filter_queryset_by_role(
        DailyCashSettlement.objects.all(),
        user,
    ).filter(
        settlement_date__gte=filter_values["start_date"],
        settlement_date__lte=filter_values["end_date"],
    ).filter(
        Q(cash_settled__gt=0) | Q(gpay_settled__gt=0),
    )
    return {
        "manual_queryset": manual_queryset,
        "settlement_queryset": settlement_queryset,
        "cash_settlement_queryset": settlement_queryset.filter(cash_settled__gt=0),
    }


def build_reconciliation_income_entries(
    user,
    filter_values,
    *,
    page_number=1,
    per_page=RECONCILIATION_HISTORY_PER_PAGE,
    include_records=True,
):
    querysets = get_reconciliation_income_querysets(user, filter_values)
    manual_queryset = querysets["manual_queryset"]
    settlement_queryset = querysets["settlement_queryset"]
    cash_settlement_queryset = querysets["cash_settlement_queryset"]

    settlement_cash_target_totals = {}
    for cash_settled_to, cash_settled_amount in cash_settlement_queryset.order_by(
        "-settlement_date",
        "-updated_at",
        "-pk",
    ).values_list("cash_settled_to", "cash_settled"):
        target_label = " ".join(str(cash_settled_to or "").split()) or "Not specified"
        settlement_cash_target_totals[target_label] = (
            settlement_cash_target_totals.get(target_label, Decimal("0.00"))
            + (cash_settled_amount or Decimal("0.00"))
        )

    manual_count = manual_queryset.count()
    settlement_count = cash_settlement_queryset.count()
    record_count = manual_count + settlement_count
    manual_income_total = _sum_amount(manual_queryset)
    manual_cash_income_total = _sum_amount(
        manual_queryset.filter(payment_method=PaymentMethod.CASH)
    )
    settlement_cash_total = _sum_field(settlement_queryset, "cash_settled")
    settlement_card_total = _sum_field(settlement_queryset, "gpay_settled")
    settlement_income_total = settlement_cash_total + settlement_card_total
    settlement_cash_target_display = ", ".join(
        f"{label} (Rs. {format_money(amount)})"
        for label, amount in settlement_cash_target_totals.items()
    ) or "-"

    resolved_page_number = 1
    entries = []
    if include_records:
        page_obj = Paginator(range(record_count), per_page).get_page(page_number)
        resolved_page_number = page_obj.number
        page_end = resolved_page_number * per_page
        page_start = max((resolved_page_number - 1) * per_page, 0)

        # Only fetch enough rows to render the requested page.
        manual_entries = [
            {
                "transaction_date": record.transaction_date,
                "title": record.title,
                "source": record.source,
                "category": record.category,
                "payment_method": record.payment_method,
                "amount": record.amount,
                "entry_type": "income",
                "entry_note": "",
                "sort_date": record.transaction_date,
                "sort_timestamp": record.created_at,
                "sort_source_priority": 1,
                "sort_pk": record.pk,
            }
            for record in manual_queryset.only(
                "transaction_date",
                "title",
                "source",
                "category",
                "payment_method",
                "amount",
                "created_at",
            ).order_by("-transaction_date", "-created_at", "-pk")[:page_end]
        ]
        settlement_entries = [
            {
                "transaction_date": settlement.settlement_date,
                "title": "Daily Cash Settlement",
                "source": (
                    f"Cash settled to {settlement.cash_settled_to}"
                    if settlement.cash_settled_to
                    else "Cash settled"
                ),
                "category": "Daily Settlement",
                "payment_method": PaymentMethod.CASH,
                "amount": settlement.cash_settled,
                "entry_type": "settlement",
                "entry_note": (
                    f"Cash settled to {' '.join(str(settlement.cash_settled_to or '').split())}"
                    if " ".join(str(settlement.cash_settled_to or "").split())
                    else "Cash settled"
                ),
                "sort_date": settlement.settlement_date,
                "sort_timestamp": settlement.updated_at,
                "sort_source_priority": 0,
                "sort_pk": settlement.pk,
            }
            for settlement in cash_settlement_queryset.only(
                "settlement_date",
                "cash_settled",
                "cash_settled_to",
                "updated_at",
            ).order_by("-settlement_date", "-updated_at", "-pk")[:page_end]
        ]
        entries = sorted(
            manual_entries + settlement_entries,
            key=lambda item: (
                item["sort_date"],
                item["sort_timestamp"],
                item["sort_source_priority"],
                item["sort_pk"],
            ),
            reverse=True,
        )[page_start:page_end]
        for entry in entries:
            entry.pop("sort_source_priority", None)
            entry.pop("sort_pk", None)

    return {
        "records": entries,
        "record_count": record_count,
        "resolved_page_number": resolved_page_number,
        "manual_income_total": manual_income_total,
        "manual_cash_income_total": manual_cash_income_total,
        "settlement_income_total": settlement_income_total,
        "settlement_cash_total": settlement_cash_total,
        "settlement_card_total": settlement_card_total,
        "settlement_cash_target_display": settlement_cash_target_display,
    }


def get_reconciliation_closing_balance_for_date(user, target_date):
    return get_reconciliation_opening_balance_for_date(
        user,
        target_date,
    ) + get_reconciliation_cash_delta_for_range(
        user,
        target_date,
        target_date,
    )


def get_first_reconciliation_activity_date(user):
    opening_balance_date = (
        filter_queryset_by_role(
            ReconciliationOpeningBalance.objects.all(),
            user,
        )
        .order_by("balance_date")
        .values_list("balance_date", flat=True)
        .first()
    )
    income_date = (
        filter_queryset_by_role(
            IncomeRecord.objects.all(),
            user,
        )
        .order_by("transaction_date")
        .values_list("transaction_date", flat=True)
        .first()
    )
    expense_date = (
        filter_queryset_by_role(
            ExpenseRecord.objects.all(),
            user,
        )
        .order_by("transaction_date")
        .values_list("transaction_date", flat=True)
        .first()
    )
    purchase_date = (
        get_purchase_base_queryset(user)
        .order_by("transaction_date", "created_at", "pk")
        .values_list("transaction_date", flat=True)
        .first()
    )
    settlement_date = (
        filter_queryset_by_role(
            DailyCashSettlement.objects.all(),
            user,
        )
        .order_by("settlement_date")
        .values_list("settlement_date", flat=True)
        .first()
    )
    candidate_dates = [
        value
        for value in (
            opening_balance_date,
            income_date,
            expense_date,
            purchase_date,
            settlement_date,
        )
        if value
    ]
    return min(candidate_dates) if candidate_dates else None


def get_saved_reconciliation_opening_balance_for_date(user, target_date):
    return (
        filter_queryset_by_role(
            ReconciliationOpeningBalance.objects.all(),
            user,
        )
        .filter(balance_date=target_date)
        .values_list("amount", flat=True)
        .first()
    )


def get_legacy_reconciliation_opening_balance_for_date(user, target_date):
    return (
        filter_queryset_by_role(
            ReconciliationIncomeEntry.objects.all(),
            user,
        )
        .filter(transaction_date=target_date, opening_balance__gt=0)
        .order_by("created_at", "pk")
        .values_list("opening_balance", flat=True)
        .first()
    )


def get_latest_reconciliation_opening_reset_before_date(user, target_date):
    saved_reset = (
        filter_queryset_by_role(
            ReconciliationOpeningBalance.objects.all(),
            user,
        )
        .filter(balance_date__lt=target_date)
        .order_by("-balance_date", "-updated_at", "-pk")
        .values("balance_date", "amount")
        .first()
    )
    legacy_reset = (
        filter_queryset_by_role(
            ReconciliationIncomeEntry.objects.all(),
            user,
        )
        .filter(transaction_date__lt=target_date, opening_balance__gt=0)
        .order_by("-transaction_date", "-created_at", "-pk")
        .values("transaction_date", "opening_balance")
        .first()
    )

    if saved_reset and legacy_reset:
        if saved_reset["balance_date"] >= legacy_reset["transaction_date"]:
            return {
                "balance_date": saved_reset["balance_date"],
                "amount": saved_reset["amount"],
            }
        return {
            "balance_date": legacy_reset["transaction_date"],
            "amount": legacy_reset["opening_balance"],
        }
    if saved_reset:
        return {
            "balance_date": saved_reset["balance_date"],
            "amount": saved_reset["amount"],
        }
    if legacy_reset:
        return {
            "balance_date": legacy_reset["transaction_date"],
            "amount": legacy_reset["opening_balance"],
        }
    return None


def get_reconciliation_purchase_non_cash_total(queryset):
    return (
        queryset.annotate(
            normalized_purchase_type=Lower(
                Trim(Coalesce("purchase_type", Value("")))
            )
        )
        .filter(
            normalized_purchase_type__in=RECONCILIATION_NON_CASH_METHOD_CASEFOLD
        )
        .aggregate(total=Sum("paid_amount"))["total"]
        or Decimal("0.00")
    )


def get_reconciliation_purchase_cash_total(queryset):
    purchase_total = _sum_field(queryset, "paid_amount")
    purchase_non_cash_total = get_reconciliation_purchase_non_cash_total(queryset)
    return purchase_total - purchase_non_cash_total


def get_reconciliation_cash_delta_for_range(user, start_date, end_date):
    if start_date is None or end_date is None or start_date > end_date:
        return Decimal("0.00")

    manual_income_total = _sum_amount(
        filter_queryset_by_role(
            IncomeRecord.objects.all(),
            user,
        ).filter(
            transaction_date__gte=start_date,
            transaction_date__lte=end_date,
        )
    )
    settlement_income_total = (
        filter_queryset_by_role(
            DailyCashSettlement.objects.all(),
            user,
        )
        .filter(
            settlement_date__gte=start_date,
            settlement_date__lte=end_date,
        )
        .aggregate(total=Sum("cash_settled") + Sum("gpay_settled"))["total"]
        or Decimal("0.00")
    )
    manual_cash_expense_total = _sum_amount(
        filter_queryset_by_role(
            ExpenseRecord.objects.all(),
            user,
        ).filter(
            transaction_date__gte=start_date,
            transaction_date__lte=end_date,
            payment_method=PaymentMethod.CASH,
        )
    )
    purchase_cash_total = get_reconciliation_purchase_cash_total(
        get_purchase_base_queryset(user).filter(
            transaction_date__gte=start_date,
            transaction_date__lte=end_date,
        )
    )
    return (
        manual_income_total
        + settlement_income_total
        - manual_cash_expense_total
        - purchase_cash_total
    )


def get_reconciliation_opening_balance_for_date(user, target_date):
    saved_balance_record = get_saved_reconciliation_opening_balance_for_date(
        user,
        target_date,
    )
    if saved_balance_record is not None:
        return saved_balance_record

    saved_opening_balance = get_legacy_reconciliation_opening_balance_for_date(
        user,
        target_date,
    )
    if saved_opening_balance is not None:
        return saved_opening_balance

    latest_reset = get_latest_reconciliation_opening_reset_before_date(
        user,
        target_date,
    )
    if latest_reset is not None:
        return latest_reset["amount"] + get_reconciliation_cash_delta_for_range(
            user,
            latest_reset["balance_date"],
            target_date - timedelta(days=1),
        )

    first_activity_date = get_first_reconciliation_activity_date(user)
    if first_activity_date is None or target_date <= first_activity_date:
        return Decimal("0.00")

    return get_reconciliation_cash_delta_for_range(
        user,
        first_activity_date,
        target_date - timedelta(days=1),
    )


def get_reconciliation_split_card_balance_for_date(user, target_date):
    settlement_card_total = (
        filter_queryset_by_role(
            DailyCashSettlement.objects.all(),
            user,
        )
        .filter(settlement_date__lte=target_date)
        .aggregate(total=Sum("gpay_settled"))["total"]
        or Decimal("0.00")
    )
    manual_card_income_total = _sum_amount(
        filter_queryset_by_role(
            IncomeRecord.objects.all(),
            user,
        ).filter(
            transaction_date__lte=target_date,
            payment_method__in=RECONCILIATION_NON_CASH_METHODS,
        )
    )
    manual_card_expense_total = _sum_amount(
        filter_queryset_by_role(
            ExpenseRecord.objects.all(),
            user,
        ).filter(
            transaction_date__lte=target_date,
            payment_method__in=RECONCILIATION_NON_CASH_METHODS,
        )
    )
    return settlement_card_total + manual_card_income_total - manual_card_expense_total


def save_reconciliation_opening_balance_for_date(user, target_date, amount):
    ReconciliationOpeningBalance.objects.update_or_create(
        user=user,
        balance_date=target_date,
        defaults={"amount": amount or Decimal("0.00")},
    )


def get_reconciliation_purchase_payment_method(purchase_type):
    normalized_purchase_type = (purchase_type or "").strip()
    if not normalized_purchase_type:
        return PaymentMethod.CASH

    if normalized_purchase_type.casefold() == PaymentMethod.CASH.casefold():
        return PaymentMethod.CASH

    if any(
        normalized_purchase_type.casefold() == method.casefold()
        for method in RECONCILIATION_NON_CASH_METHODS
    ):
        return normalized_purchase_type

    return PaymentMethod.CASH


def get_reconciliation_expense_querysets(user, filter_values):
    manual_queryset = filter_queryset_by_role(
        ExpenseRecord.objects.all(),
        user,
    ).select_related("supplier").filter(
        transaction_date__gte=filter_values["start_date"],
        transaction_date__lte=filter_values["end_date"],
    )
    purchase_queryset = get_purchase_base_queryset(user).filter(
        transaction_date__gte=filter_values["start_date"],
        transaction_date__lte=filter_values["end_date"],
    )
    return {
        "manual_queryset": manual_queryset,
        "purchase_queryset": purchase_queryset,
    }


def build_reconciliation_expense_entries(
    user,
    filter_values,
    *,
    page_number=1,
    per_page=RECONCILIATION_HISTORY_PER_PAGE,
    include_records=True,
):
    querysets = get_reconciliation_expense_querysets(user, filter_values)
    manual_queryset = querysets["manual_queryset"]
    purchase_queryset = querysets["purchase_queryset"]

    manual_count = manual_queryset.count()
    purchase_count = purchase_queryset.count()
    record_count = manual_count + purchase_count
    manual_expense_total = _sum_amount(manual_queryset)
    manual_cash_expense_total = _sum_amount(
        manual_queryset.filter(payment_method=PaymentMethod.CASH)
    )
    manual_non_cash_expense_total = _sum_amount(
        manual_queryset.filter(payment_method__in=RECONCILIATION_NON_CASH_METHODS)
    )
    purchase_total = _sum_field(purchase_queryset, "paid_amount")
    purchase_non_cash_total = get_reconciliation_purchase_non_cash_total(
        purchase_queryset
    )
    purchase_cash_total = purchase_total - purchase_non_cash_total

    resolved_page_number = 1
    entries = []
    if include_records:
        page_obj = Paginator(range(record_count), per_page).get_page(page_number)
        resolved_page_number = page_obj.number
        page_end = resolved_page_number * per_page
        page_start = max((resolved_page_number - 1) * per_page, 0)

        manual_entries = [
            {
                "transaction_date": record.transaction_date,
                "title": record.title,
                "vendor": record.supplier_display,
                "category": record.category,
                "payment_method": record.payment_method,
                "amount": record.amount,
                "entry_type": "expense",
                "sort_date": record.transaction_date,
                "sort_timestamp": record.created_at,
                "sort_source_priority": 1,
                "sort_pk": record.pk,
            }
            for record in manual_queryset.only(
                "transaction_date",
                "title",
                "vendor",
                "category",
                "payment_method",
                "amount",
                "created_at",
                "supplier__name",
            ).order_by("-transaction_date", "-created_at", "-pk")[:page_end]
        ]
        purchase_entries = []
        for purchase in purchase_queryset.order_by(
            "-transaction_date",
            "-created_at",
            "-pk",
        )[:page_end]:
            purchase_payment_method = get_reconciliation_purchase_payment_method(
                purchase.purchase_type
            )
            purchase_entries.append(
                {
                    "transaction_date": purchase.transaction_date,
                    "title": purchase.invoice_number or "Purchase Record",
                    "vendor": purchase.supplier_name or (
                        purchase.supplier.name if purchase.supplier_id else "-"
                    ),
                    "category": purchase.purchase_type or "Purchase",
                    "payment_method": purchase_payment_method,
                    "amount": purchase.paid_amount,
                    "entry_type": "purchase",
                    "sort_date": purchase.transaction_date,
                    "sort_timestamp": purchase.created_at,
                    "sort_source_priority": 0,
                    "sort_pk": purchase.pk,
                }
            )

        entries = sorted(
            manual_entries + purchase_entries,
            key=lambda item: (
                item["sort_date"],
                item["sort_timestamp"],
                item["sort_source_priority"],
                item["sort_pk"],
            ),
            reverse=True,
        )[page_start:page_end]
        for entry in entries:
            entry.pop("sort_source_priority", None)
            entry.pop("sort_pk", None)

    return {
        "records": entries,
        "record_count": record_count,
        "resolved_page_number": resolved_page_number,
        "manual_expense_total": manual_expense_total,
        "manual_cash_expense_total": manual_cash_expense_total,
        "manual_non_cash_expense_total": manual_non_cash_expense_total,
        "purchase_cash_total": purchase_cash_total,
        "purchase_non_cash_total": purchase_non_cash_total,
        "purchase_total": purchase_total,
    }


def build_ranked_category_overview(rows):
    category_rows = list(rows)
    highest_total = max(
        [row["total"] or Decimal("0.00") for row in category_rows],
        default=Decimal("1.00"),
    )
    if highest_total == 0:
        highest_total = Decimal("1.00")

    aggregate_total = sum(
        (row["total"] or Decimal("0.00") for row in category_rows),
        Decimal("0.00"),
    )
    overview = []
    for index, row in enumerate(category_rows, start=1):
        total = row["total"] or Decimal("0.00")
        overview.append(
            {
                "rank": index,
                "category": row["category"] or "Uncategorized",
                "total": total,
                "width": round((total / highest_total) * 100, 2),
                "share": round((total / aggregate_total) * 100, 1)
                if aggregate_total
                else Decimal("0.0"),
            }
        )
    return overview


def get_next_supplier_code():
    last_supplier = Supplier.objects.order_by("-pk").first()
    next_number = 1 if last_supplier is None else last_supplier.pk + 1
    return f"SUP-{next_number:04d}"


def build_page_url(request, page_number, page_param="page", **extra_params):
    query_params = request.GET.copy()
    query_params[page_param] = page_number
    for key, value in extra_params.items():
        if value is None:
            query_params.pop(key, None)
        else:
            query_params[key] = value
    return f"{request.path}?{query_params.urlencode()}"


def paginate_record_list(
    request,
    records,
    *,
    page_param="page",
    per_page=20,
    history_view=None,
):
    paginator = Paginator(records, per_page)
    page_obj = paginator.get_page((request.GET.get(page_param) or "").strip() or 1)
    url_overrides = {}
    if history_view:
        url_overrides["view"] = history_view

    previous_page_url = ""
    next_page_url = ""
    if page_obj.has_previous():
        previous_page_url = build_page_url(
            request,
            page_obj.previous_page_number(),
            page_param=page_param,
            **url_overrides,
        )
    if page_obj.has_next():
        next_page_url = build_page_url(
            request,
            page_obj.next_page_number(),
            page_param=page_param,
            **url_overrides,
        )

    return {
        "records": list(page_obj.object_list),
        "page_obj": page_obj,
        "previous_page_url": previous_page_url,
        "next_page_url": next_page_url,
        "pagination_links": [
            {
                "number": page_number,
                "url": build_page_url(
                    request,
                    page_number,
                    page_param=page_param,
                    **url_overrides,
                ),
                "is_current": page_number == page_obj.number,
            }
            for page_number in page_obj.paginator.page_range
        ],
        "showing_from": page_obj.start_index(),
        "showing_to": page_obj.end_index(),
    }


def get_safe_next_url(request, fallback_url):
    next_url = (request.GET.get("next") or request.POST.get("next") or "").strip()
    if next_url and url_has_allowed_host_and_scheme(
        next_url,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        return next_url
    return fallback_url


def get_selected_date(request, parameter_name="as_of_date"):
    selected_date = parse_date((request.GET.get(parameter_name) or "").strip())
    return selected_date or date.today()


def get_purchase_base_queryset(user):
    queryset = filter_queryset_by_role(PurchaseRecord.objects.all(), user)
    return (
        queryset.filter(
            Q(source_reference__startswith=MANUAL_PURCHASE_SOURCE_PREFIX)
            | Q(source_reference__startswith=SQLSERVER_PURCHASE_SOURCE_PREFIX)
            | Q(source_reference="")
        )
        .select_related("supplier", "user")
    )


def get_raw_purchase_filter_values(params):
    has_pending_value = (params.get("has_pending") or "").strip()
    return {
        "search": (params.get("search") or "").strip(),
        "supplier_name": (params.get("supplier_name") or "").strip(),
        "invoice_number": (params.get("invoice_number") or "").strip(),
        "saved_by": (params.get("saved_by") or "").strip(),
        "pending_amount": (params.get("pending_amount") or "").strip(),
        "has_pending": "1" if has_pending_value else "",
        "date_from": (params.get("date_from") or "").strip(),
        "date_to": (params.get("date_to") or "").strip(),
    }


def get_purchase_filter_values(params):
    filter_values = get_raw_purchase_filter_values(params)
    if not any(filter_values.values()):
        today_value = date.today().isoformat()
        filter_values["date_from"] = today_value
        filter_values["date_to"] = today_value
    elif filter_values["date_from"] and not filter_values["date_to"]:
        filter_values["date_to"] = filter_values["date_from"]
    elif filter_values["date_to"] and not filter_values["date_from"]:
        filter_values["date_from"] = filter_values["date_to"]
    return filter_values


def get_purchase_redirect_params(params):
    redirect_params = {}
    for key in (
        "search",
        "supplier_name",
        "invoice_number",
        "saved_by",
        "pending_amount",
        "has_pending",
        "date_from",
        "date_to",
    ):
        value = (params.get(key) or "").strip()
        if value:
            redirect_params[key] = value
    return redirect_params


def build_purchase_redirect_url(params):
    redirect_params = get_purchase_redirect_params(params)
    if not redirect_params:
        return reverse("purchase-list")
    return f"{reverse('purchase-list')}?{urlencode(redirect_params)}"


def apply_purchase_filters(queryset, filter_values):
    search = filter_values["search"]
    supplier_name = filter_values["supplier_name"]
    invoice_number = filter_values["invoice_number"]
    saved_by = filter_values["saved_by"]
    pending_amount = filter_values["pending_amount"]
    has_pending = filter_values["has_pending"]
    date_from = parse_date(filter_values["date_from"])
    date_to = parse_date(filter_values["date_to"])

    if search:
        queryset = queryset.filter(
            Q(supplier_name__icontains=search)
            | Q(invoice_number__icontains=search)
            | Q(user__username__icontains=search)
        )
    if supplier_name:
        queryset = queryset.filter(supplier_name__icontains=supplier_name)
    if invoice_number:
        queryset = queryset.filter(invoice_number__icontains=invoice_number)
    if saved_by:
        queryset = queryset.filter(user__username__icontains=saved_by)
    if pending_amount:
        try:
            queryset = queryset.filter(pending_amount=Decimal(pending_amount))
        except InvalidOperation:
            if has_pending:
                queryset = queryset.filter(pending_amount__gt=Decimal("0.00"))
    elif has_pending:
        queryset = queryset.filter(pending_amount__gt=Decimal("0.00"))
    if date_from and date_to and date_from > date_to:
        date_from, date_to = date_to, date_from
    if date_from:
        queryset = queryset.filter(transaction_date__gte=date_from)
    if date_to:
        queryset = queryset.filter(transaction_date__lte=date_to)
    return queryset


def summarize_purchase_queryset(queryset):
    aggregates = queryset.aggregate(
        total_amount=Sum("total_amount"),
        paid_amount=Sum("paid_amount"),
        pending_amount=Sum("pending_amount"),
    )
    return {
        "count": queryset.count(),
        "total_amount": aggregates["total_amount"] or Decimal("0.00"),
        "paid_amount": aggregates["paid_amount"] or Decimal("0.00"),
        "pending_amount": aggregates["pending_amount"] or Decimal("0.00"),
    }


def get_raw_sales_filter_values(params):
    return {
        "bill_no": (params.get("bill_no") or "").strip(),
        "customer_name": (params.get("customer_name") or "").strip(),
        "payment_mode": (params.get("payment_mode") or "").strip(),
        "date_from": (params.get("date_from") or "").strip(),
        "date_to": (params.get("date_to") or "").strip(),
    }


def get_sales_filter_values(params):
    filter_values = get_raw_sales_filter_values(params)
    if not filter_values["date_from"] and not filter_values["date_to"]:
        today_value = date.today().isoformat()
        filter_values["date_from"] = today_value
        filter_values["date_to"] = today_value
    elif filter_values["date_from"] and not filter_values["date_to"]:
        filter_values["date_to"] = filter_values["date_from"]
    elif filter_values["date_to"] and not filter_values["date_from"]:
        filter_values["date_from"] = filter_values["date_to"]
    return filter_values


def apply_sales_filters(queryset, filter_values):
    bill_no = filter_values["bill_no"]
    customer_name = filter_values["customer_name"]
    payment_mode = filter_values["payment_mode"]
    date_from = parse_date(filter_values["date_from"])
    date_to = parse_date(filter_values["date_to"])

    if bill_no:
        queryset = queryset.filter(bill_no__icontains=bill_no)
    if customer_name:
        queryset = queryset.filter(customer_name__icontains=customer_name)
    if payment_mode == SPLIT_PAYMENT_MODE_FILTER:
        queryset = queryset.filter(
            split_cash_amount__gt=Decimal("0.00"),
            split_card_amount__gt=Decimal("0.00"),
        )
    elif payment_mode and payment_mode != "All":
        queryset = queryset.filter(payment_mode=payment_mode)
    if date_from and date_to and date_from > date_to:
        date_from, date_to = date_to, date_from
    if date_from:
        queryset = queryset.filter(sale_date__gte=date_from)
    if date_to:
        queryset = queryset.filter(sale_date__lte=date_to)
    return queryset


def summarize_sales_records(records):
    records = list(records)
    split_records = [
        record for record in records if record.split_cash_amount > 0 or record.split_card_amount > 0
    ]
    credit_records = [
        record for record in records if get_effective_sales_balance_amount(record) > 0
    ]
    return {
        "count": len(records),
        "total_amount": sum(
            (record.net_amount for record in records),
            Decimal("0.00"),
        ),
        "received_amount": sum(
            (record.effective_received_amount for record in records),
            Decimal("0.00"),
        ),
        "balance_amount": sum(
            (record.effective_balance_amount for record in records),
            Decimal("0.00"),
        ),
        "split_cash_amount": sum(
            (record.split_cash_amount for record in records),
            Decimal("0.00"),
        ),
        "split_card_amount": sum(
            (record.split_card_amount for record in records),
            Decimal("0.00"),
        ),
        "split_bill_count": len(split_records),
        "split_total_amount": sum(
            (record.split_total_amount for record in split_records),
            Decimal("0.00"),
        ),
        "credit_count": len(credit_records),
        "credit_total": sum(
            (record.effective_balance_amount for record in credit_records),
            Decimal("0.00"),
        ),
    }


def get_sales_report_base_queryset(user):
    return SalesLedgerRecord.objects.filter(
        is_cancelled=False,
        sale_date__isnull=False,
    )


def summarize_sales_queryset_for_range(queryset, start_date, end_date):
    filtered_queryset = queryset.filter(
        sale_date__gte=start_date,
        sale_date__lte=end_date,
    )
    aggregates = filtered_queryset.aggregate(total_amount=Sum("net_amount"))
    return {
        "count": filtered_queryset.count(),
        "total_amount": aggregates["total_amount"] or Decimal("0.00"),
    }


def get_sales_report_filter_values(params):
    today = date.today()
    range_key = (params.get("range") or "").strip().lower() or "month"
    raw_start_date = (params.get("start_date") or "").strip()
    raw_end_date = (params.get("end_date") or "").strip()
    payment_mode = (params.get("payment_mode") or "").strip()
    start_date = parse_date(raw_start_date)
    end_date = parse_date(raw_end_date)

    preset_labels = {
        "today": "Today",
        "week": "This week",
        "month": "This month",
        "last30": "Last 30 days",
        "custom": "Custom range",
    }

    if range_key == "today":
        start_date = today
        end_date = today
    elif range_key == "week":
        start_date = today - timedelta(days=today.weekday())
        end_date = today
    elif range_key == "last30":
        start_date = today - timedelta(days=29)
        end_date = today
    elif range_key == "custom":
        if start_date is None and end_date is None:
            start_date = today.replace(day=1)
            end_date = today
        elif start_date is None:
            start_date = end_date
        elif end_date is None:
            end_date = start_date
    else:
        range_key = "month"
        start_date = today.replace(day=1)
        end_date = today

    if start_date is None:
        start_date = today.replace(day=1)
    if end_date is None:
        end_date = today
    if start_date > end_date:
        start_date, end_date = end_date, start_date

    scope_label = (
        f"{preset_labels.get(range_key, 'Filtered range')} "
        f"({start_date:%d %b %Y} to {end_date:%d %b %Y})"
    )
    if range_key == "custom":
        scope_label = f"Custom range ({start_date:%d %b %Y} to {end_date:%d %b %Y})"

    return {
        "range": range_key,
        "start_date": start_date,
        "end_date": end_date,
        "payment_mode": payment_mode,
        "scope_label": scope_label,
        "range_label": preset_labels.get(range_key, "Filtered range"),
        "has_active_filters": range_key != "month" or bool(payment_mode),
    }


def apply_sales_report_filters(queryset, filter_values):
    queryset = queryset.filter(
        sale_date__gte=filter_values["start_date"],
        sale_date__lte=filter_values["end_date"],
    )
    payment_mode = filter_values["payment_mode"]
    if payment_mode == SPLIT_PAYMENT_MODE_FILTER:
        queryset = queryset.filter(
            split_cash_amount__gt=Decimal("0.00"),
            split_card_amount__gt=Decimal("0.00"),
        )
    elif payment_mode:
        queryset = queryset.filter(payment_mode=payment_mode)
    return queryset


def build_sales_chart_rows(queryset, period, max_points=12):
    if period == "day":
        chart_title = "Daily sales"
        base_rows = [
            {
                "bucket": row["sale_date"],
                "label": row["sale_date"].strftime("%d %b %Y"),
                "short_label": row["sale_date"].strftime("%d %b"),
                "total": row["total"] or Decimal("0.00"),
                "count": row["count"] or 0,
            }
            for row in (
                queryset.values("sale_date")
                .annotate(total=Sum("net_amount"), count=Count("source_sale_no"))
                .order_by("sale_date")
            )
        ]
    elif period == "week":
        chart_title = "Weekly sales"
        base_rows = []
        for row in (
            queryset.annotate(bucket=TruncWeek("sale_date"))
            .values("bucket")
            .annotate(total=Sum("net_amount"), count=Count("source_sale_no"))
            .order_by("bucket")
        ):
            bucket_value = row["bucket"]
            if hasattr(bucket_value, "date"):
                bucket_value = bucket_value.date()
            base_rows.append(
                {
                    "bucket": bucket_value,
                    "label": f"Week of {bucket_value:%d %b %Y}",
                    "short_label": bucket_value.strftime("%d %b"),
                    "total": row["total"] or Decimal("0.00"),
                    "count": row["count"] or 0,
                }
            )
    else:
        chart_title = "Monthly sales"
        base_rows = []
        for row in (
            queryset.annotate(bucket=TruncMonth("sale_date"))
            .values("bucket")
            .annotate(total=Sum("net_amount"), count=Count("source_sale_no"))
            .order_by("bucket")
        ):
            bucket_value = row["bucket"]
            if hasattr(bucket_value, "date"):
                bucket_value = bucket_value.date()
            base_rows.append(
                {
                    "bucket": bucket_value,
                    "label": bucket_value.strftime("%B %Y"),
                    "short_label": bucket_value.strftime("%b %Y"),
                    "total": row["total"] or Decimal("0.00"),
                    "count": row["count"] or 0,
                }
            )

    rows = base_rows[-max_points:]
    return finalize_sales_chart_rows(
        rows,
        title=chart_title,
        scope_label=f"Filtered range | showing {len(rows)} bucket{'' if len(rows) == 1 else 's'}",
    )


def finalize_sales_chart_rows(rows, title, scope_label):
    highest_total = max(
        [row["total"] for row in rows],
        default=Decimal("1.00"),
    )
    if highest_total == 0:
        highest_total = Decimal("1.00")

    for row in rows:
        row["height"] = round((row["total"] / highest_total) * 100, 2)

    return {
        "title": title,
        "scope_label": scope_label,
        "rows": rows,
        "total_sales": sum((row["total"] for row in rows), Decimal("0.00")),
        "total_bills": sum((row["count"] for row in rows), 0),
    }


def build_sales_payment_split_rows(records):
    cash_total = Decimal("0.00")
    upi_total = Decimal("0.00")
    credit_total = Decimal("0.00")

    for record in records:
        if record.has_manual_split:
            cash_total += record.split_cash_amount
            upi_total += record.split_card_amount
            credit_total += max(
                record.net_amount - record.split_total_amount,
                Decimal("0.00"),
            )
            continue

        if record.payment_mode == SalesPaymentMode.CARD:
            upi_total += record.effective_received_amount
        else:
            cash_total += record.effective_received_amount

        credit_total += record.effective_balance_amount

    split_rows = [
        {"key": "cash", "label": "Cash", "amount": cash_total},
        {"key": "upi", "label": "UPI / Card", "amount": upi_total},
        {"key": "credit", "label": "Credit", "amount": credit_total},
    ]
    split_total = sum((row["amount"] for row in split_rows), Decimal("0.00"))
    for row in split_rows:
        row["share"] = round((row["amount"] / split_total) * 100, 1) if split_total else Decimal("0.0")
        row["width"] = row["share"]
    return split_rows, split_total


def build_sales_report_workspace(user, params, include_export_records=False):
    today = date.today()
    base_queryset = get_sales_report_base_queryset(user)
    filter_values = get_sales_report_filter_values(params)
    filtered_queryset = apply_sales_report_filters(base_queryset, filter_values)
    split_records = list(
        filtered_queryset.only(
            "bill_no",
            "sale_date",
            "customer_name",
            "net_amount",
            "received_amount",
            "balance_amount",
            "split_cash_amount",
            "split_card_amount",
            "payment_mode",
        ).order_by("sale_date", "source_sale_no")
    )
    filtered_summary = summarize_sales_records(split_records)
    payment_split_rows, payment_split_total = build_sales_payment_split_rows(split_records)
    payment_split_lookup = {row["key"]: row["amount"] for row in payment_split_rows}
    total_bills = filtered_summary["count"]
    average_bill_amount = (
        round(filtered_summary["total_amount"] / total_bills, 2)
        if total_bills
        else Decimal("0.00")
    )
    workspace = {
        "today_summary": summarize_sales_queryset_for_range(base_queryset, today, today),
        "week_summary": summarize_sales_queryset_for_range(
            base_queryset,
            today - timedelta(days=today.weekday()),
            today,
        ),
        "month_summary": summarize_sales_queryset_for_range(
            base_queryset,
            today.replace(day=1),
            today,
        ),
        "sales_report_filters": filter_values,
        "filtered_summary": filtered_summary,
        "average_bill_amount": average_bill_amount,
        "sales_charts": [
            build_sales_chart_rows(filtered_queryset, period="day", max_points=14),
            build_sales_chart_rows(filtered_queryset, period="week", max_points=12),
            build_sales_chart_rows(filtered_queryset, period="month", max_points=12),
        ],
        "payment_split_rows": payment_split_rows,
        "payment_split_total": payment_split_total,
        "payment_split_lookup": payment_split_lookup,
        "payment_split_scope_label": filter_values["scope_label"],
        "payment_mode_options": [
            *SalesPaymentMode.choices[:-1],
            (SPLIT_PAYMENT_MODE_FILTER, SPLIT_PAYMENT_MODE_FILTER),
            SalesPaymentMode.choices[-1],
        ],
        "top_selling_products": [],
        "category_sales": [],
    }
    if include_export_records:
        workspace["export_records"] = list(
            filtered_queryset.order_by("-sale_date", "-source_sale_no")
        )
    return workspace


def get_sales_redirect_params(params):
    redirect_params = {}
    for key in ("bill_no", "customer_name", "payment_mode", "date_from", "date_to"):
        value = (params.get(key) or "").strip()
        if value:
            redirect_params[key] = value
    return redirect_params


def build_sales_redirect_url(params):
    redirect_params = get_sales_redirect_params(params)
    if not redirect_params:
        return reverse("sales-list")
    return f"{reverse('sales-list')}?{urlencode(redirect_params)}"

def get_raw_expense_filter_values(request):
    return {
        "start_date": (request.GET.get("start_date") or "").strip(),
        "end_date": (request.GET.get("end_date") or "").strip(),
        "category": (request.GET.get("category") or "").strip(),
        "payment_method": (request.GET.get("payment_method") or "").strip(),
    }


def get_expense_filter_values(request):
    filter_values = get_raw_expense_filter_values(request)
    if not any(filter_values.values()):
        today_value = date.today().isoformat()
        filter_values["start_date"] = today_value
        filter_values["end_date"] = today_value
        return filter_values

    if filter_values["start_date"] and not filter_values["end_date"]:
        filter_values["end_date"] = filter_values["start_date"]
    if filter_values["end_date"] and not filter_values["start_date"]:
        filter_values["start_date"] = filter_values["end_date"]
    return filter_values


def apply_expense_filters(queryset, filter_values):
    start_date = parse_date(filter_values["start_date"])
    end_date = parse_date(filter_values["end_date"])
    category = filter_values["category"]
    payment_method = filter_values["payment_method"]

    if start_date and end_date and start_date > end_date:
        start_date, end_date = end_date, start_date
    if start_date:
        queryset = queryset.filter(transaction_date__gte=start_date)
    if end_date:
        queryset = queryset.filter(transaction_date__lte=end_date)
    if category and category != "All":
        queryset = queryset.filter(category__iexact=category)
    if payment_method and payment_method != "All":
        queryset = queryset.filter(payment_method=payment_method)
    return queryset


def get_expense_category_options(user):
    return get_saved_expense_category_options(user)


def build_settlement_redirect_url(settlement_date):
    return (
        f"{reverse('daily-settlement')}?"
        f"{urlencode({'selected_date': settlement_date.isoformat(), 'entry_date': settlement_date.isoformat()})}"
    )


def get_expense_redirect_params(params):
    redirect_params = {}
    for key in ("start_date", "end_date", "category", "payment_method"):
        value = (params.get(key) or "").strip()
        if value:
            redirect_params[key] = value
    return redirect_params


def build_expense_redirect_url(params):
    redirect_params = get_expense_redirect_params(params)
    if not redirect_params:
        return reverse("expense-list")
    return f"{reverse('expense-list')}?{urlencode(redirect_params)}"


def parse_money_value(value):
    if isinstance(value, Decimal):
        return value
    if value in (None, ""):
        return Decimal("0.00")
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0.00")


CASH_DENOMINATION_VALUES = (2000, 500, 200, 100, 50, 20, 10, 5, 2, 1)


def normalize_count_value(value):
    try:
        count = int(str(value or "0").strip())
    except (TypeError, ValueError):
        return 0
    return max(count, 0)


def parse_cash_denominations_payload(raw_payload):
    if not raw_payload:
        return {}
    try:
        parsed_payload = json.loads(raw_payload)
    except (TypeError, ValueError):
        return {}
    if not isinstance(parsed_payload, dict):
        return {}

    normalized = {}
    for denomination in CASH_DENOMINATION_VALUES:
        count = normalize_count_value(parsed_payload.get(str(denomination)))
        if count > 0:
            normalized[str(denomination)] = count
    return normalized


def build_cash_denominations_payload(denominations):
    normalized = {
        str(denomination): normalize_count_value(denominations.get(str(denomination)))
        for denomination in CASH_DENOMINATION_VALUES
        if normalize_count_value(denominations.get(str(denomination))) > 0
    }
    return json.dumps(normalized) if normalized else ""


def get_cash_denominations_total(denominations):
    total = Decimal("0.00")
    for denomination in CASH_DENOMINATION_VALUES:
        count = normalize_count_value(denominations.get(str(denomination)))
        if count:
            total += Decimal(str(denomination)) * Decimal(str(count))
    return total


def get_cash_difference_amount(cash_denomination_total, cash_in_hand):
    return parse_money_value(cash_denomination_total) - parse_money_value(cash_in_hand)


def build_cash_denomination_rows(denominations):
    rows = []
    for denomination in CASH_DENOMINATION_VALUES:
        count = normalize_count_value(denominations.get(str(denomination)))
        rows.append(
            {
                "value": denomination,
                "count": count,
                "total": Decimal(str(denomination)) * Decimal(str(count)),
            }
        )
    return rows


def get_effective_sales_payment_amount(record):
    return record.effective_received_amount


def get_effective_sales_balance_amount(record):
    return record.effective_balance_amount


def get_credit_bill_records_for_date(settlement_date):
    records = list(
        SalesLedgerRecord.objects.filter(
            is_cancelled=False,
            sale_date=settlement_date,
        ).order_by("-sale_date", "-source_sale_no")
    )
    credit_records = [
        record for record in records if get_effective_sales_balance_amount(record) > 0
    ]
    return sorted(
        credit_records,
        key=lambda record: (-record.effective_balance_amount, record.bill_no),
    )


def build_sales_settlement_summary(settlement_date):
    sales_queryset = SalesLedgerRecord.objects.filter(
        is_cancelled=False,
        sale_date=settlement_date,
    ).order_by("source_sale_no")
    sales_count = sales_queryset.count()

    if not sales_count:
        return {
            "gpay_settled": Decimal("0.00"),
            "cash_settled": Decimal("0.00"),
            "sales_count": 0,
            "manual_split_count": 0,
        }

    gpay_total = Decimal("0.00")
    cash_total = Decimal("0.00")
    manual_split_count = 0

    for record in sales_queryset:
        if record.has_manual_split:
            manual_split_count += 1
            cash_total += record.split_cash_amount
            gpay_total += record.split_card_amount
            continue

        effective_amount = get_effective_sales_payment_amount(record)
        if record.payment_mode == SalesPaymentMode.CARD:
            gpay_total += effective_amount
        elif record.payment_mode == SalesPaymentMode.CASH:
            cash_total += effective_amount

    return {
        "gpay_settled": gpay_total,
        "cash_settled": cash_total,
        "sales_count": sales_count,
        "manual_split_count": manual_split_count,
    }
    

def get_settlement_filter_values(params):
    selected_date = parse_date(
        (params.get("selected_date") or params.get("entry_date") or params.get("settlement_date") or "").strip()
    ) or date.today()
    return {
        "selected_date": selected_date,
    }


def get_settlement_entry_date(params, default_date):
    entry_date = parse_date(
        (params.get("entry_date") or params.get("settlement_date") or "").strip()
    )
    return entry_date or default_date


def get_default_settlement_opening_balance(user, settlement_date):
    previous_settlement = (
        filter_queryset_by_role(DailyCashSettlement.objects.all(), user).filter(
            settlement_date__lt=settlement_date,
        )
        .order_by("-settlement_date", "-updated_at", "-pk")
        .first()
    )
    if previous_settlement:
        return previous_settlement.closing_balance

    income_total = _sum_amount(
        filter_queryset_by_role(IncomeRecord.objects.all(), user).exclude(
            category__iexact=INCOME_CATEGORY_OFFICE,
        ).filter(
            transaction_date__lt=settlement_date,
        )
    )
    expense_total = _sum_amount(
        filter_queryset_by_role(ExpenseRecord.objects.all(), user).filter(
            transaction_date__lt=settlement_date,
        )
    )
    return income_total - expense_total


def get_counter_income_total_for_date(user, settlement_date, payment_methods=None):
    queryset = filter_queryset_by_role(IncomeRecord.objects.all(), user).filter(
        transaction_date=settlement_date,
        category__iexact=INCOME_CATEGORY_COUNTER,
    )
    return _sum_amount(_apply_payment_method_filter(queryset, payment_methods))


def get_counter_expense_total_for_date(user, settlement_date, payment_methods=None):
    queryset = filter_queryset_by_role(ExpenseRecord.objects.all(), user).filter(
        transaction_date=settlement_date,
        category__iexact=COUNTER_EXPENSE_CATEGORY,
    )
    return _sum_amount(_apply_payment_method_filter(queryset, payment_methods))


def get_income_total_for_date(user, settlement_date, payment_methods=None):
    queryset = filter_queryset_by_role(IncomeRecord.objects.all(), user).filter(
        transaction_date=settlement_date,
    )
    return _sum_amount(_apply_payment_method_filter(queryset, payment_methods))


def get_expense_total_for_date(user, settlement_date, payment_methods=None):
    queryset = filter_queryset_by_role(ExpenseRecord.objects.all(), user).filter(
        transaction_date=settlement_date,
    )
    return _sum_amount(_apply_payment_method_filter(queryset, payment_methods))


def build_settlement_payment_summary(user, settlement_date):
    cash_income_amount = get_counter_income_total_for_date(
        user,
        settlement_date,
        PaymentMethod.CASH,
    )
    non_cash_counter_income_amount = get_counter_income_total_for_date(
        user,
        settlement_date,
        RECONCILIATION_NON_CASH_METHODS,
    )
    cash_expense_amount = get_counter_expense_total_for_date(
        user,
        settlement_date,
        PaymentMethod.CASH,
    )
    non_cash_expense_amount = get_expense_total_for_date(
        user,
        settlement_date,
        SETTLEMENT_NON_CASH_EXPENSE_METHODS,
    )
    return {
        "cash_income_amount": cash_income_amount,
        "non_cash_counter_income_amount": non_cash_counter_income_amount,
        "cash_expense_amount": cash_expense_amount,
        "non_cash_expense_amount": non_cash_expense_amount,
    }


def build_settlement_autofill_summary(user, settlement_date):
    previous_settlement = (
        filter_queryset_by_role(DailyCashSettlement.objects.all(), user).filter(
            settlement_date__lt=settlement_date,
        )
        .order_by("-settlement_date", "-updated_at", "-pk")
        .first()
    )
    sales_summary = build_sales_settlement_summary(settlement_date)
    counter_expense_queryset = filter_queryset_by_role(
        ExpenseRecord.objects.all(),
        user,
    ).filter(
        transaction_date=settlement_date,
        category__iexact=COUNTER_EXPENSE_CATEGORY,
    )
    payment_summary = build_settlement_payment_summary(user, settlement_date)
    if previous_settlement:
        opening_balance = previous_settlement.closing_balance
    else:
        opening_balance = get_default_settlement_opening_balance(user, settlement_date)

    sales_ledger_cash = sales_summary["cash_settled"]
    counter_income_amount = payment_summary["cash_income_amount"]
    gpay_income_amount = payment_summary["non_cash_counter_income_amount"]
    gpay_settled = sales_summary["gpay_settled"]
    gpay_expense_amount = Decimal("0.00")
    gpay_net_amount = gpay_settled + gpay_income_amount
    if sales_summary["sales_count"] > 0:
        settlement_source = "sales"
    else:
        settlement_source = "manual"

    expense_amount = payment_summary["cash_expense_amount"]
    cash_in_hand = get_cash_in_hand_amount(
        opening_balance,
        sales_ledger_cash,
        expense_amount,
        counter_income_amount=counter_income_amount,
    )

    return {
        "opening_balance": opening_balance,
        "sales_ledger_cash": sales_ledger_cash,
        "counter_income_amount": counter_income_amount,
        "gpay_income_amount": gpay_income_amount,
        "gpay_settled": gpay_settled,
        "gpay_expense_amount": gpay_expense_amount,
        "gpay_net_amount": gpay_net_amount,
        "cash_settled": Decimal("0.00"),
        "cash_in_hand": cash_in_hand,
        "expense_amount": expense_amount,
        "upi_count": 0,
        "cash_count": 0,
        "expense_count": counter_expense_queryset.count(),
        "sales_count": sales_summary["sales_count"],
        "manual_split_count": sales_summary["manual_split_count"],
        "settlement_source": settlement_source,
        "previous_settlement_date": (
            previous_settlement.settlement_date if previous_settlement else None
        ),
        "previous_closing_balance": (
            previous_settlement.closing_balance if previous_settlement else opening_balance
        ),
    }


def build_settlement_preview(values, cash_denomination_total=Decimal("0.00")):
    opening_balance = parse_money_value(values.get("opening_balance"))
    sales_ledger_cash = parse_money_value(values.get("sales_ledger_cash"))
    counter_income_amount = parse_money_value(values.get("counter_income_amount"))
    gpay_settled = parse_money_value(values.get("gpay_settled"))
    cash_settled = parse_money_value(values.get("cash_settled"))
    expense_amount = parse_money_value(values.get("expense_amount"))
    closing_balance = get_settlement_closing_balance(
        opening_balance,
        sales_ledger_cash,
        expense_amount,
        cash_settled,
        cash_denomination_total,
        counter_income_amount=counter_income_amount,
    )
    cash_in_hand = get_cash_in_hand_amount(
        opening_balance,
        sales_ledger_cash,
        expense_amount,
        counter_income_amount=counter_income_amount,
    )
    total_amount = opening_balance + sales_ledger_cash + gpay_settled
    actual_sales = total_amount - opening_balance
    return {
        "opening_balance": opening_balance,
        "sales_ledger_cash": sales_ledger_cash,
        "counter_income_amount": counter_income_amount,
        "gpay_settled": gpay_settled,
        "cash_settled": cash_settled,
        "cash_in_hand": cash_in_hand,
        "expense_amount": expense_amount,
        "closing_balance": closing_balance,
        "total_amount": total_amount,
        "actual_sales": actual_sales,
    }


def build_invoice_filename(invoice_number):
    safe_invoice = re.sub(r"[^A-Za-z0-9_-]+", "_", invoice_number or "").strip("_")
    if not safe_invoice:
        safe_invoice = "purchase"
    return f"{safe_invoice}_invoice.pdf"


def build_purchase_invoice_response(purchase):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=0)
    page_width, page_height = A4
    saved_on = purchase.get_saved_on_display()
    payment_history = build_purchase_payment_history(purchase)
    pdf.setTitle(f"Purchase Invoice - {purchase.invoice_number}")

    premium_navy = colors.HexColor("#12263A")
    premium_gold = colors.HexColor("#D5B26F")
    premium_cream = colors.HexColor("#F7F1E3")
    panel_fill = colors.HexColor("#F9F7F2")
    accent_fill = colors.HexColor("#E8D8B5")
    text_dark = colors.HexColor("#22313F")
    muted_text = colors.HexColor("#667786")
    line_color = colors.HexColor("#D9C7A0")

    def draw_label_value(x_position, y_position, label, value):
        pdf.setFillColor(muted_text)
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(x_position, y_position, label.upper())
        pdf.setFillColor(text_dark)
        pdf.setFont("Helvetica", 11)
        pdf.drawString(x_position, y_position - 14, str(value))

    def draw_amount_card(x_position, y_position, width, title, amount, fill_color):
        pdf.setFillColor(fill_color)
        pdf.roundRect(x_position, y_position, width, 62, 10, fill=1, stroke=0)
        pdf.setFillColor(text_dark)
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(x_position + 14, y_position + 42, title)
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(x_position + 14, y_position + 20, f"Rs. {amount}")

    def draw_footer(message):
        pdf.setStrokeColor(premium_gold)
        pdf.line(32, 60, page_width - 32, 60)
        pdf.setFillColor(muted_text)
        pdf.setFont("Helvetica-Oblique", 9)
        pdf.drawString(32, 44, message)

    def draw_payment_history_page_header():
        pdf.showPage()
        pdf.setFillColor(premium_navy)
        pdf.roundRect(32, page_height - 126, page_width - 64, 86, 18, fill=1, stroke=0)
        pdf.setFillColor(premium_cream)
        pdf.setFont("Helvetica-Bold", 20)
        pdf.drawString(52, page_height - 78, "Payment Tracking")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(52, page_height - 98, "Full invoice payment register with date, time, amount, and staff entry.")
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawRightString(page_width - 52, page_height - 76, f"Invoice No: {purchase.invoice_number}")
        pdf.setFillColor(text_dark)
        pdf.setFont("Helvetica", 10)
        pdf.drawString(52, page_height - 152, f"Supplier: {purchase.supplier_name}")
        pdf.drawString(52, page_height - 170, f"Purchase Date: {purchase.transaction_date or '-'}")
        pdf.drawString(280, page_height - 152, f"Total Amount: Rs. {purchase.total_amount}")
        pdf.drawString(280, page_height - 170, f"Pending Amount: Rs. {purchase.pending_amount}")
        pdf.setStrokeColor(line_color)
        pdf.line(32, page_height - 190, page_width - 32, page_height - 190)
        pdf.setFillColor(muted_text)
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawString(40, page_height - 210, "DATE")
        pdf.drawString(102, page_height - 210, "TIME")
        pdf.drawString(160, page_height - 210, "ENTRY")
        pdf.drawString(296, page_height - 210, "AMOUNT")
        pdf.drawString(370, page_height - 210, "STAFF")
        pdf.drawRightString(page_width - 40, page_height - 210, "BALANCE")
        pdf.setStrokeColor(line_color)
        pdf.line(32, page_height - 218, page_width - 32, page_height - 218)
        return page_height - 238

    pdf.setFillColor(premium_navy)
    pdf.roundRect(32, page_height - 142, page_width - 64, 102, 18, fill=1, stroke=0)
    pdf.setFillColor(premium_gold)
    pdf.circle(72, page_height - 91, 18, fill=1, stroke=0)
    pdf.setFillColor(premium_navy)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawCentredString(72, page_height - 96, "M")
    pdf.setFillColor(premium_cream)
    pdf.setFont("Helvetica-Bold", 22)
    pdf.drawString(104, page_height - 78, "Mahilmart Purchase Invoice")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(104, page_height - 98, "Official supplier purchase invoice for Mahilmart")
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawRightString(page_width - 52, page_height - 72, f"Invoice No: {purchase.invoice_number}")
    pdf.setFont("Helvetica", 10)
    pdf.drawRightString(page_width - 52, page_height - 92, f"Saved On: {saved_on}")

    pdf.setFillColor(panel_fill)
    pdf.roundRect(32, page_height - 332, page_width - 64, 160, 16, fill=1, stroke=0)
    pdf.setStrokeColor(line_color)
    pdf.roundRect(32, page_height - 332, page_width - 64, 160, 16, fill=0, stroke=1)
    pdf.setFillColor(text_dark)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(52, page_height - 196, "Purchase Details")

    draw_label_value(52, page_height - 220, "Supplier Name", purchase.supplier_name)
    draw_label_value(280, page_height - 220, "Type", purchase.purchase_type)
    draw_label_value(52, page_height - 262, "Invoice Number", purchase.invoice_number)
    draw_label_value(280, page_height - 262, "Prepared By", purchase.user.username)
    draw_label_value(52, page_height - 304, "Purchase Date", purchase.transaction_date or "-")
    draw_label_value(280, page_height - 304, "Saved On", saved_on)
    draw_label_value(52, page_height - 346, "Supplier Code", purchase.supplier.supplier_code if purchase.supplier_id else "-")
    draw_label_value(280, page_height - 346, "Uploaded File", "Attached" if purchase.attachment else "No Attachment")

    pdf.setFillColor(text_dark)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(32, page_height - 402, "Amount Summary")
    draw_amount_card(32, page_height - 482, 165, "Total Amount", purchase.total_amount, accent_fill)
    draw_amount_card(214, page_height - 482, 165, "Paid Amount", purchase.paid_amount, colors.HexColor("#D9F0DF"))
    draw_amount_card(396, page_height - 482, 165, "Pending Amount", purchase.pending_amount, colors.HexColor("#F7D6CF"))

    pdf.setFillColor(panel_fill)
    pdf.roundRect(32, page_height - 624, page_width - 64, 108, 16, fill=1, stroke=0)
    pdf.setStrokeColor(line_color)
    pdf.roundRect(32, page_height - 624, page_width - 64, 108, 16, fill=0, stroke=1)
    pdf.setFillColor(text_dark)
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawString(52, page_height - 544, "Notes")
    pdf.setFont("Helvetica", 10)
    notes = purchase.notes.strip() or "No notes added for this purchase."
    note_y = page_height - 564
    for line in textwrap.wrap(notes, width=92):
        pdf.drawString(52, note_y, line)
        note_y -= 14
        if note_y < page_height - 608:
            break

    signature_y = 100
    pdf.setStrokeColor(line_color)
    pdf.line(60, signature_y, 230, signature_y)
    pdf.line(340, signature_y, 510, signature_y)
    pdf.setFillColor(text_dark)
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(60, signature_y - 16, "Supplier Signature")
    pdf.drawString(340, signature_y - 16, "Authorized Sign")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(340, signature_y - 32, f"Prepared By: {purchase.user.username}")

    draw_footer("This invoice was generated from the Mahilmart purchase module.")

    detail_y = draw_payment_history_page_header()
    for entry in payment_history:
        note_lines = textwrap.wrap(
            entry["notes"] or "No note added for this entry.",
            width=50,
        )[:2]
        row_height = 32 + (len(note_lines) * 11)

        if detail_y - row_height < 92:
            draw_footer("Payment tracking continued from the Mahilmart purchase module.")
            detail_y = draw_payment_history_page_header()

        pdf.setFillColor(text_dark)
        pdf.setFont("Helvetica", 9)
        pdf.drawString(40, detail_y, entry["payment_date"])
        pdf.drawString(102, detail_y, entry["recorded_time"])
        pdf.drawString(160, detail_y, entry["label"])
        pdf.drawString(296, detail_y, f"Rs. {entry['amount']}")
        pdf.drawString(370, detail_y, entry["recorded_by"][:18])
        pdf.drawRightString(page_width - 40, detail_y, f"Rs. {entry['running_pending']}")

        pdf.setFillColor(muted_text)
        pdf.setFont("Helvetica", 8.5)
        pdf.drawString(160, detail_y - 13, f"Paid Till This Entry: Rs. {entry['running_paid']}")

        note_y = detail_y - 26
        for line in note_lines:
            pdf.drawString(160, note_y, line)
            note_y -= 11

        pdf.setStrokeColor(line_color)
        pdf.line(32, detail_y - row_height, page_width - 32, detail_y - row_height)
        detail_y -= row_height + 12

    draw_footer("This invoice includes date-wise payment history recorded by Mahilmart staff.")
    pdf.save()

    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="{build_invoice_filename(purchase.invoice_number)}"'
    )
    return response


def format_money(amount):
    return f"{amount:.2f}"


def get_daily_settlement_email_recipients():
    raw_value = getattr(settings, "CONTACT_RECEIVER_EMAIL", "")
    if isinstance(raw_value, str):
        return [email.strip() for email in raw_value.split(",") if email.strip()]
    if raw_value:
        return [str(email).strip() for email in raw_value if str(email).strip()]
    return []


def build_cash_denomination_summary(denominations):
    parts = []
    for denomination in CASH_DENOMINATION_VALUES:
        count = normalize_count_value(denominations.get(str(denomination)))
        if count:
            parts.append(f"{denomination} x {count}")
    return ", ".join(parts) if parts else "-"


def build_daily_settlement_email_context(settlement):
    payment_summary = build_settlement_payment_summary(
        settlement.user,
        settlement.settlement_date,
    )
    counter_income_amount = payment_summary["cash_income_amount"]
    gpay_income_amount = payment_summary["non_cash_counter_income_amount"]
    gpay_expense_amount = Decimal("0.00")
    gpay_net_amount = settlement.gpay_settled + gpay_income_amount
    sales_cash_amount = get_sales_cash_from_settlement(settlement)
    cash_denominations = settlement.cash_denominations or {}
    preview_values = {
        "opening_balance": settlement.opening_balance,
        "sales_ledger_cash": sales_cash_amount,
        "counter_income_amount": counter_income_amount,
        "gpay_settled": settlement.gpay_settled,
        "cash_settled": settlement.cash_settled,
        "expense_amount": settlement.expense_amount,
    }
    preview = build_settlement_preview(
        preview_values,
        cash_denomination_total=settlement.cash_denomination_total,
    )
    credit_bill_records = get_credit_bill_records_for_date(settlement.settlement_date)[:10]
    credit_bill_total = sum(
        (record.effective_balance_amount for record in credit_bill_records),
        Decimal("0.00"),
    )
    saved_on = (
        timezone.localtime(settlement.updated_at).strftime("%d-%m-%Y %I:%M %p")
        if settlement.updated_at
        else "-"
    )
    return {
        "settlement": settlement,
        "saved_by": settlement.user.username,
        "saved_on": saved_on,
        "preview": preview,
        "counter_income_amount": counter_income_amount,
        "gpay_income_amount": gpay_income_amount,
        "gpay_expense_amount": gpay_expense_amount,
        "gpay_net_amount": gpay_net_amount,
        "sales_cash_amount": sales_cash_amount,
        "cash_denomination_rows": build_cash_denomination_rows(cash_denominations),
        "cash_denomination_summary": build_cash_denomination_summary(cash_denominations),
        "credit_bill_records": credit_bill_records,
        "credit_bill_count": len(credit_bill_records),
        "credit_bill_total": credit_bill_total,
    }


def build_daily_settlement_email_subject(settlement):
    return f"Daily Settlement Details - {settlement.settlement_date:%d-%m-%Y}"


def build_daily_settlement_email_body(settlement):
    return render_to_string(
        "tracker/emails/daily_settlement_email.txt",
        build_daily_settlement_email_context(settlement),
    )


def build_daily_settlement_email_html(settlement):
    return render_to_string(
        "tracker/emails/daily_settlement_email.html",
        build_daily_settlement_email_context(settlement),
    )


def send_daily_settlement_email(settlement):
    recipients = get_daily_settlement_email_recipients()
    if not recipients:
        return []

    send_mail(
        subject=build_daily_settlement_email_subject(settlement),
        message=build_daily_settlement_email_body(settlement),
        html_message=build_daily_settlement_email_html(settlement),
        from_email=(getattr(settings, "DEFAULT_FROM_EMAIL", "") or None),
        recipient_list=recipients,
        fail_silently=False,
    )
    return recipients


def create_purchase_payment(
    purchase,
    user,
    amount,
    notes="",
    payment_date=None,
    update_totals=True,
):
    payment = PurchasePayment.objects.create(
        purchase=purchase,
        user=user,
        amount=amount,
        payment_date=payment_date or date.today(),
        notes=notes.strip(),
    )
    if update_totals:
        purchase.paid_amount += amount
        purchase.save(update_fields=["paid_amount", "pending_amount", "updated_at"])
    return payment


def build_purchase_payment_history(purchase):
    history = []
    tracked_total = Decimal("0.00")
    running_paid = Decimal("0.00")
    purchase_date = purchase.transaction_date or purchase.created_at.date()

    history.append(
        (
            purchase_date,
            purchase.created_at,
            {
                "entry_type": "purchase",
                "label": "Purchase Created",
                "amount": format_money(purchase.total_amount),
                "payment_date": purchase_date.strftime("%d-%m-%Y"),
                "recorded_on": purchase.created_at.strftime("%d-%m-%Y %I:%M %p"),
                "recorded_time": purchase.created_at.strftime("%I:%M %p"),
                "recorded_by": purchase.user.username,
                "notes": purchase.notes.strip() or "Purchase record created.",
                "running_paid": format_money(running_paid),
                "running_pending": format_money(purchase.total_amount - running_paid),
            },
        )
    )

    for payment in purchase.payments.select_related("user").order_by(
        "payment_date",
        "created_at",
        "pk",
    ):
        tracked_total += payment.amount
        running_paid += payment.amount
        history.append(
            (
                payment.payment_date,
                payment.created_at,
                {
                    "entry_type": "payment",
                    "label": "Payment Entry",
                    "amount": format_money(payment.amount),
                    "payment_date": payment.payment_date.strftime("%d-%m-%Y"),
                    "recorded_on": payment.created_at.strftime("%d-%m-%Y %I:%M %p"),
                    "recorded_time": payment.created_at.strftime("%I:%M %p"),
                    "recorded_by": payment.user.username,
                    "notes": payment.notes.strip(),
                    "running_paid": format_money(running_paid),
                    "running_pending": format_money(purchase.total_amount - running_paid),
                },
            )
        )

    legacy_balance = purchase.paid_amount - tracked_total
    if legacy_balance > 0:
        running_paid += legacy_balance
        history.append(
            (
                purchase_date,
                purchase.created_at,
                {
                    "entry_type": "opening",
                    "label": "Opening Paid Amount",
                    "amount": format_money(legacy_balance),
                    "payment_date": purchase_date.strftime("%d-%m-%Y"),
                    "recorded_on": purchase.created_at.strftime("%d-%m-%Y %I:%M %p"),
                    "recorded_time": purchase.created_at.strftime("%I:%M %p"),
                    "recorded_by": purchase.user.username,
                    "notes": "Saved with the original purchase record.",
                    "running_paid": format_money(running_paid),
                    "running_pending": format_money(purchase.total_amount - running_paid),
                },
            )
        )

    history.sort(key=lambda item: (item[0], item[1]))
    return [item[2] for item in history]


def build_purchase_detail_payload(purchase):
    payment_history = build_purchase_payment_history(purchase)
    payment_entries = [
        entry for entry in payment_history if entry["entry_type"] in {"payment", "opening"}
    ]
    return {
        "supplier_name": purchase.supplier_name,
        "purchase_type": purchase.purchase_type,
        "invoice_number": purchase.invoice_number,
        "saved_by": purchase.user.username,
        "saved_on": purchase.get_saved_on_display(),
        "purchase_date": purchase.transaction_date.strftime("%d-%m-%Y") if purchase.transaction_date else "-",
        "total_amount": format_money(purchase.total_amount),
        "paid_amount": format_money(purchase.paid_amount),
        "pending_amount": format_money(purchase.pending_amount),
        "last_payment_date": payment_entries[-1]["payment_date"] if payment_entries else "-",
        "invoice_url": reverse("purchase-invoice", args=[purchase.pk]),
        "attachment_url": purchase.attachment.url if purchase.attachment else "",
        "pay_url": reverse("purchase-pay", args=[purchase.pk]),
        "can_pay": purchase.pending_amount > 0,
        "payment_history": payment_history,
    }


class HomeRedirectView(RedirectView):
    pattern_name = "dashboard"

    def get_redirect_url(self, *args, **kwargs):
        if self.request.user.is_authenticated:
            return reverse_lazy(
                get_first_accessible_route_name(self.request.user) or "access-denied"
            )
        return reverse_lazy("login")


class AccessDeniedView(LoginRequiredMixin, TemplateView):
    template_name = "tracker/access_denied.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        fallback_route_name = get_first_accessible_route_name(self.request.user)
        fallback_label = ""
        if fallback_route_name:
            fallback_label = next(
                (
                    label
                    for label, _field_name, route_name in TRACKER_PERMISSION_ITEMS
                    if route_name == fallback_route_name
                ),
                "",
            )
        context.update(
            {
                "fallback_route_name": fallback_route_name,
                "fallback_url": (
                    reverse(fallback_route_name) if fallback_route_name else ""
                ),
                "fallback_label": fallback_label,
            }
        )
        return context


class DashboardView(ModulePermissionRequiredMixin, TemplateView):
    permission_field = "allow_dashboard"
    permission_denied_message = "You do not have access to Dashboard."
    template_name = "tracker/dashboard.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        selected_date = get_selected_date(self.request)
        previous_date = selected_date - timedelta(days=1)
        income_queryset = filter_queryset_by_role(IncomeRecord.objects.all(), user)
        expense_queryset = filter_queryset_by_role(
            ExpenseRecord.objects.all(),
            user,
        ).select_related("supplier")
        settlement_queryset = filter_queryset_by_role(
            DailyCashSettlement.objects.all(),
            user,
        )
        settlement_record = settlement_queryset.filter(settlement_date=selected_date).first()
        income_total = _sum_settlement_income(
            settlement_queryset.filter(settlement_date__lte=selected_date)
        )
        expense_total = _sum_amount(expense_queryset.filter(transaction_date__lte=selected_date))
        opening_balance = get_default_settlement_opening_balance(user, selected_date)
        selected_income = _sum_settlement_income(
            settlement_queryset.filter(settlement_date=selected_date)
        )
        selected_expense = _sum_amount(expense_queryset.filter(transaction_date=selected_date))
        closing_balance = Decimal("0.00")
        if settlement_record:
            opening_balance = settlement_record.opening_balance
            closing_balance = settlement_record.closing_balance
        current_month_income = _sum_settlement_income(
            settlement_queryset.filter(
                settlement_date__year=selected_date.year,
                settlement_date__month=selected_date.month,
                settlement_date__lte=selected_date,
            )
        )
        current_month_expense = _sum_amount(
            expense_queryset.filter(
                transaction_date__year=selected_date.year,
                transaction_date__month=selected_date.month,
                transaction_date__lte=selected_date,
            )
        )
        selected_day_balance = Decimal("0.00")
        if settlement_record:
            selected_day_balance = settlement_record.actual_sales
        balance = income_total - expense_total
        savings_rate = 0
        if income_total:
            savings_rate = round(float((balance / income_total) * 100), 1)
        savings_rate_meter = min(max(savings_rate, 0), 100)

        context.update(
            {
                "income_total": income_total,
                "expense_total": expense_total,
                "balance": balance,
                "opening_balance": opening_balance,
                "closing_balance": closing_balance,
                "dashboard_settlement": settlement_record,
                "selected_income": selected_income,
                "selected_expense": selected_expense,
                "selected_day_balance": selected_day_balance,
                "selected_date": selected_date,
                "previous_date": previous_date,
                "current_month_income": current_month_income,
                "current_month_expense": current_month_expense,
                "savings_rate": savings_rate,
                "savings_rate_meter": savings_rate_meter,
                "recent_income": income_queryset.filter(transaction_date__lte=selected_date)[:5],
                "recent_expenses": expense_queryset.filter(transaction_date__lte=selected_date)[:5],
                "monthly_overview": build_monthly_overview(user, anchor_date=selected_date),
                "supplier_count": filter_queryset_by_role(Supplier.objects.all(), user).count(),
            }
        )
        return context


class AutoLoadPaginatedListView(LoginRequiredMixin, ListView):
    paginate_by = 50

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        page_obj = context.get("page_obj")
        context["next_page_url"] = ""
        context["visible_count"] = len(context.get(self.context_object_name, []))
        if page_obj and page_obj.has_next():
            context["next_page_url"] = build_page_url(
                self.request, page_obj.next_page_number()
            )
        return context


class AdminRequiredMixin(LoginRequiredMixin):
    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            return self.handle_no_permission()
        if not request.user.is_superuser:
            messages.error(request, "Only admin users can manage users.")
            return redirect(get_first_accessible_route_name(request.user) or "access-denied")
        return super().dispatch(request, *args, **kwargs)


class IncomeListView(ModulePermissionRequiredMixin, AutoLoadPaginatedListView):
    permission_field = "allow_income"
    permission_denied_message = "You do not have access to Income."
    model = IncomeRecord
    template_name = "tracker/income_list.html"
    context_object_name = "records"
    paginate_by = 10

    def get_base_entries(self):
        return build_income_ledger_entries(self.request.user)

    def get_queryset(self):
        return apply_income_filters(
            self.get_base_entries(),
            get_income_filter_values(self.request),
        )

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        all_records = self.get_base_entries()
        filtered_records = self.get_queryset()
        today = date.today()
        raw_filter_values = get_raw_income_filter_values(self.request)
        filter_values = get_income_filter_values(self.request)
        month_records = [
            record
            for record in all_records
            if record.get("transaction_date")
            and record["transaction_date"].year == today.year
            and record["transaction_date"].month == today.month
        ]
        today_records = [
            record for record in all_records if record.get("transaction_date") == today
        ]
        context["income_filters"] = filter_values
        context["has_active_filters"] = any(raw_filter_values.values())
        context["is_default_today_view"] = not context["has_active_filters"]
        context["month_summary"] = summarize_income_entries(month_records)
        context["today_summary"] = summarize_income_entries(today_records)
        context["filtered_summary"] = summarize_income_entries(filtered_records)
        context["page_total"] = context["filtered_summary"]["total"]
        context["page_count"] = len(filtered_records)
        context["category_options"] = get_income_category_options(
            all_records,
            selected_category=filter_values["category"],
        )
        context["payment_method_options"] = get_income_payment_method_options(
            all_records,
            selected_payment_method=filter_values["payment_method"],
        )
        context["current_url"] = self.request.get_full_path()
        return context


class IncomeCreateView(ModulePermissionRequiredMixin, CreateView):
    permission_field = "allow_income"
    permission_denied_message = "You do not have access to Income."
    model = IncomeRecord
    form_class = IncomeForm
    template_name = "tracker/income_form.html"
    success_url = reverse_lazy("income-list")

    def get_return_url(self):
        return get_safe_next_url(self.request, reverse("income-list"))

    def get_success_url(self):
        return self.get_return_url()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.user = self.request.user
        messages.success(self.request, "Income record created successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        purpose_map = get_income_category_purpose_map(self.request.user)
        form_category = (
            normalize_income_category_name(
                context["form"].data.get(context["form"].add_prefix("category"))
                if context["form"].is_bound
                else getattr(context["form"].instance, "category", "")
            )
            or DEFAULT_INCOME_CATEGORY
        )
        context["income_category_purpose_map"] = purpose_map
        context["purpose_form"] = IncomePurposeForm(
            initial={"category": form_category}
        )
        context["default_income_purpose_count"] = len(
            purpose_map.get(form_category, [])
        )
        context["default_income_category"] = form_category
        context["next_url"] = self.get_return_url()
        context["form_page_url"] = self.request.get_full_path()
        context["is_editing"] = False
        return context


class IncomePurposeCreateView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_income"
    permission_denied_message = "You do not have access to Income."

    def get_next_url(self):
        next_url = (
            self.request.GET.get("next") or self.request.POST.get("next") or ""
        ).strip()
        return next_url or reverse("income-add")

    def post(self, request, *args, **kwargs):
        form = IncomePurposeForm(request.POST)
        is_ajax_request = request.headers.get("x-requested-with") == "XMLHttpRequest"
        if form.is_valid():
            category_name = normalize_income_category_name(
                form.cleaned_data["category"]
            )
            purpose_name = normalize_income_purpose_name(form.cleaned_data["name"])
            existing = (
                IncomePurpose.objects.filter(user=request.user)
                .filter(category__iexact=category_name, name__iexact=purpose_name)
                .first()
            )
            if existing is not None:
                created = False
                saved_purpose = existing
            else:
                saved_purpose = get_or_create_role_income_purpose(
                    request.user,
                    category_name,
                    purpose_name,
                )
                created = True

            purpose_map = get_income_category_purpose_map(request.user)
            category_purposes = purpose_map.get(category_name, [])

            if is_ajax_request:
                return JsonResponse(
                    {
                        "ok": True,
                        "created": created,
                        "name": saved_purpose.name,
                        "category": saved_purpose.category,
                        "purpose_count": len(category_purposes),
                        "purpose_options": category_purposes,
                        "message": (
                            "Income purpose created successfully."
                            if created
                            else (
                                f"{saved_purpose.name} already exists for "
                                f"{saved_purpose.category}."
                            )
                        ),
                    }
                )

            if created:
                messages.success(request, "Income purpose created successfully.")
            else:
                messages.info(
                    request,
                    (
                        f"{saved_purpose.name} already exists for "
                        f"{saved_purpose.category}."
                    ),
                )
            return redirect(self.get_next_url())

        if is_ajax_request:
            return JsonResponse(
                {
                    "ok": False,
                    "errors": {
                        field_name: [
                            error["message"]
                            for error in field_errors
                        ]
                        for field_name, field_errors in form.errors.get_json_data().items()
                    },
                },
                status=400,
            )

        first_error = next(
            iter(
                next(iter(form.errors.values()), ["Unable to save the purpose right now."])
            ),
            "Unable to save the purpose right now.",
        )
        messages.error(request, first_error)
        return redirect(self.get_next_url())


class IncomeUpdateView(ModulePermissionRequiredMixin, UpdateView):
    permission_field = "allow_income"
    permission_denied_message = "You do not have access to Income."
    model = IncomeRecord
    form_class = IncomeForm
    template_name = "tracker/income_form.html"
    success_url = reverse_lazy("income-list")

    def get_queryset(self):
        return filter_queryset_by_role(IncomeRecord.objects.all(), self.request.user)

    def get_return_url(self):
        return get_safe_next_url(self.request, reverse("income-list"))

    def get_success_url(self):
        return self.get_return_url()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        return kwargs

    def form_valid(self, form):
        messages.success(self.request, "Income record updated successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        purpose_map = get_income_category_purpose_map(self.request.user)
        form_category = (
            normalize_income_category_name(
                context["form"].data.get(context["form"].add_prefix("category"))
                if context["form"].is_bound
                else getattr(context["form"].instance, "category", "")
            )
            or DEFAULT_INCOME_CATEGORY
        )
        context["income_category_purpose_map"] = purpose_map
        context["purpose_form"] = IncomePurposeForm(
            initial={"category": form_category}
        )
        context["default_income_purpose_count"] = len(
            purpose_map.get(form_category, [])
        )
        context["default_income_category"] = form_category
        context["next_url"] = self.get_return_url()
        context["form_page_url"] = self.request.get_full_path()
        context["is_editing"] = True
        return context


class IncomeDeleteView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_income"
    permission_denied_message = "You do not have access to Income."

    def get_queryset(self):
        return filter_queryset_by_role(IncomeRecord.objects.all(), self.request.user)

    def post(self, request, *args, **kwargs):
        income_record = get_object_or_404(self.get_queryset(), pk=kwargs["pk"])
        income_record.delete()
        messages.success(request, "Income record deleted successfully.")
        return redirect(get_safe_next_url(request, reverse("income-list")))


class ExpenseListView(ModulePermissionRequiredMixin, AutoLoadPaginatedListView):
    permission_field = "allow_expenses"
    permission_denied_message = "You do not have access to Expenses."
    model = ExpenseRecord
    template_name = "tracker/expense_list.html"
    context_object_name = "records"

    def get_base_queryset(self):
        return (
            filter_queryset_by_role(ExpenseRecord.objects.all(), self.request.user)
            .select_related("supplier")
            .order_by("-transaction_date", "-created_at", "-pk")
        )

    def get_queryset(self):
        return apply_expense_filters(
            self.get_base_queryset(),
            get_expense_filter_values(self.request),
        )

    def get_edit_record(self):
        record_id = (self.request.GET.get("edit") or "").strip()
        if not record_id.isdigit():
            return None
        return self.get_base_queryset().filter(pk=record_id).first()

    def build_entry_rows(self):
        edit_record = self.get_edit_record()
        if edit_record:
            return [
                {
                    "id": str(edit_record.pk),
                    "transaction_date": edit_record.transaction_date.isoformat(),
                    "category": edit_record.category,
                    "purpose": edit_record.title,
                    "amount": format_money(edit_record.amount),
                    "payment_method": edit_record.payment_method,
                }
            ]
        return [
            {
                "id": "",
                "transaction_date": date.today().isoformat(),
                "category": COUNTER_EXPENSE_CATEGORY,
                "purpose": "",
                "amount": "",
                "payment_method": PaymentMethod.CASH,
            }
        ]

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip()
        if action == "delete":
            return self.handle_delete(request)
        return self.handle_save(request)

    def handle_delete(self, request):
        record_id = (request.POST.get("record_id") or "").strip()
        record = self.get_base_queryset().filter(pk=record_id).first()
        if record is None:
            messages.error(request, "Expense record not found.")
            return redirect(build_expense_redirect_url(request.POST))

        record.delete()
        messages.success(request, "Expense record deleted successfully.")
        return redirect(build_expense_redirect_url(request.POST))

    def handle_save(self, request):
        record_ids = request.POST.getlist("entry_id")
        transaction_dates = request.POST.getlist("entry_date")
        categories = request.POST.getlist("entry_category")
        purposes = request.POST.getlist("entry_purpose")
        amounts = request.POST.getlist("entry_amount")
        payment_methods = request.POST.getlist("entry_payment_method")

        row_count = max(
            len(record_ids),
            len(transaction_dates),
            len(categories),
            len(purposes),
            len(amounts),
            len(payment_methods),
        )
        created_count = 0
        updated_count = 0
        skipped_count = 0

        valid_payment_methods = {value for value, _label in PaymentMethod.choices}

        for index in range(row_count):
            record_id = (record_ids[index] if index < len(record_ids) else "").strip()
            transaction_date_value = (
                transaction_dates[index] if index < len(transaction_dates) else ""
            ).strip()
            category_value = normalize_expense_category_name(
                categories[index] if index < len(categories) else ""
            )
            purpose_value = (purposes[index] if index < len(purposes) else "").strip()
            amount_value = (amounts[index] if index < len(amounts) else "").strip()
            payment_method_value = (
                payment_methods[index] if index < len(payment_methods) else ""
            ).strip()

            if not any(
                [
                    record_id,
                    transaction_date_value,
                    category_value,
                    purpose_value,
                    amount_value,
                    payment_method_value,
                ]
            ):
                continue

            transaction_date = parse_date(transaction_date_value)
            amount = parse_money_value(amount_value)

            if (
                not transaction_date
                or not category_value
                or not purpose_value
                or amount <= 0
                or payment_method_value not in valid_payment_methods
            ):
                skipped_count += 1
                continue

            if record_id.isdigit():
                record = self.get_base_queryset().filter(pk=record_id).first()
                if record is None:
                    skipped_count += 1
                    continue
                updated_count += 1
            else:
                record = ExpenseRecord(user=request.user)
                created_count += 1

            record.title = purpose_value[:120]
            record.vendor = purpose_value[:120]
            record.category = category_value[:80]
            record.amount = amount
            record.transaction_date = transaction_date
            record.payment_method = payment_method_value
            record.save()
            get_or_create_role_expense_category(request.user, category_value)

        if created_count or updated_count:
            message_bits = []
            if created_count:
                message_bits.append(
                    f"{created_count} expense record{'s' if created_count != 1 else ''} added"
                )
            if updated_count:
                message_bits.append(
                    f"{updated_count} expense record{'s' if updated_count != 1 else ''} updated"
                )
            messages.success(request, ", ".join(message_bits) + ".")
        else:
            messages.error(request, "Enter at least one complete expense row to save.")

        if skipped_count:
            messages.warning(
                request,
                f"{skipped_count} incomplete row{'s were' if skipped_count != 1 else ' was'} skipped.",
            )

        return redirect(build_expense_redirect_url(request.POST))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        all_records = self.get_base_queryset()
        filtered_records = self.get_queryset()
        today = date.today()
        raw_filter_values = get_raw_expense_filter_values(self.request)
        filter_values = get_expense_filter_values(self.request)
        context["page_total"] = _sum_amount(filtered_records)
        context["page_count"] = filtered_records.count()
        context["supplier_count"] = filter_queryset_by_role(
            Supplier.objects.all(),
            user,
        ).count()
        context["expense_filters"] = filter_values
        context["has_active_filters"] = any(raw_filter_values.values())
        context["is_default_today_view"] = not context["has_active_filters"]
        context["month_total"] = _sum_amount(
            all_records.filter(
                transaction_date__year=today.year,
                transaction_date__month=today.month,
            )
        )
        context["today_total"] = _sum_amount(all_records.filter(transaction_date=today))
        context["filtered_total"] = context["page_total"]
        context["category_options"] = get_expense_category_options(user)
        context["expense_category_purpose_map"] = get_expense_category_purpose_map(user)
        context["payment_method_options"] = PaymentMethod.choices
        context["entry_rows"] = self.build_entry_rows()
        context["editing_record"] = self.get_edit_record()
        context["filter_query"] = urlencode(get_expense_redirect_params(self.request.GET))
        return context


class ExpenseCreateView(ModulePermissionRequiredMixin, CreateView):
    permission_field = "allow_expenses"
    permission_denied_message = "You do not have access to Expenses."
    model = ExpenseRecord
    form_class = ExpenseForm
    template_name = "tracker/expense_form.html"
    success_url = reverse_lazy("expense-list")

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.user = self.request.user
        messages.success(self.request, "Expense record created successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        purpose_map = get_expense_category_purpose_map(self.request.user)
        form_category = (
            normalize_expense_category_name(
                context["form"].data.get(context["form"].add_prefix("category"))
                if context["form"].is_bound
                else getattr(context["form"].instance, "category", "")
            )
            or COUNTER_EXPENSE_CATEGORY
        )
        context["supplier_count"] = filter_queryset_by_role(
            Supplier.objects.all(),
            self.request.user,
        ).count()
        context["counter_expense_category"] = COUNTER_EXPENSE_CATEGORY
        context["office_expense_category"] = OFFICE_EXPENSE_CATEGORY
        context["expense_category_purpose_map"] = purpose_map
        context["purpose_form"] = ExpensePurposeForm(
            initial={"category": form_category}
        )
        context["default_expense_category"] = form_category
        context["default_expense_purpose_count"] = len(
            purpose_map.get(form_category, [])
        )
        return context


class ExpenseCategoryListView(ModulePermissionRequiredMixin, TemplateView):
    permission_field = "allow_expenses"
    permission_denied_message = "You do not have access to Expenses."
    template_name = "tracker/expense_category_list.html"

    def get_next_url(self):
        next_url = (self.request.GET.get("next") or self.request.POST.get("next") or "").strip()
        return next_url or reverse("expense-add")

    def post(self, request, *args, **kwargs):
        form = ExpenseCategoryForm(request.POST)
        is_ajax_request = request.headers.get("x-requested-with") == "XMLHttpRequest"
        if form.is_valid():
            category_name = normalize_expense_category_name(form.cleaned_data["name"])
            existing = (
                ExpenseCategory.objects.filter(user=request.user)
                .filter(name__iexact=category_name)
                .first()
            )
            if existing is not None:
                if is_ajax_request:
                    return JsonResponse(
                        {
                            "ok": True,
                            "created": False,
                            "name": existing.name,
                            "category_count": len(
                                get_expense_category_options(request.user)
                            ),
                            "message": (
                                f"{existing.name} already exists for this account."
                            ),
                        }
                    )
                messages.info(request, f"{existing.name} already exists for this account.")
            else:
                ExpenseCategory.objects.create(
                    user=request.user,
                    name=category_name,
                )
                if is_ajax_request:
                    return JsonResponse(
                        {
                            "ok": True,
                            "created": True,
                            "name": category_name,
                            "category_count": len(
                                get_expense_category_options(request.user)
                            ),
                            "message": "Expense category created successfully.",
                        }
                    )
                messages.success(request, "Expense category created successfully.")
            return redirect(self.get_next_url())
        if is_ajax_request:
            return JsonResponse(
                {
                    "ok": False,
                    "errors": {
                        field_name: [
                            error["message"]
                            for error in field_errors
                        ]
                        for field_name, field_errors in form.errors.get_json_data().items()
                    },
                },
                status=400,
            )
        context = self.get_context_data(form=form)
        return self.render_to_response(context)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        categories = ensure_expense_categories_for_role(self.request.user)
        usage_lookup = {
            normalize_expense_category_name(item["category"]).casefold(): item["total"]
            for item in (
                filter_queryset_by_role(ExpenseRecord.objects.all(), self.request.user)
                .exclude(category="")
                .values("category")
                .annotate(total=Count("id"))
            )
        }
        context["form"] = kwargs.get("form") or ExpenseCategoryForm()
        context["category_records"] = [
            {
                "name": category.name,
                "usage_count": usage_lookup.get(category.name.casefold(), 0),
            }
            for category in categories
        ]
        context["category_count"] = len(context["category_records"])
        context["next_url"] = self.get_next_url()
        return context


class ExpensePurposeCreateView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_expenses"
    permission_denied_message = "You do not have access to Expenses."

    def get_next_url(self):
        next_url = (
            self.request.GET.get("next") or self.request.POST.get("next") or ""
        ).strip()
        return next_url or reverse("expense-add")

    def post(self, request, *args, **kwargs):
        form = ExpensePurposeForm(request.POST)
        is_ajax_request = request.headers.get("x-requested-with") == "XMLHttpRequest"
        if form.is_valid():
            category_name = normalize_expense_category_name(
                form.cleaned_data["category"]
            )
            purpose_name = normalize_expense_purpose_name(form.cleaned_data["name"])
            existing = (
                ExpensePurpose.objects.filter(user=request.user)
                .filter(category__iexact=category_name, name__iexact=purpose_name)
                .first()
            )
            if existing is not None:
                created = False
                saved_purpose = existing
            else:
                saved_purpose = get_or_create_role_expense_purpose(
                    request.user,
                    category_name,
                    purpose_name,
                )
                created = True

            purpose_map = get_expense_category_purpose_map(request.user)
            category_purposes = purpose_map.get(category_name, [])

            if is_ajax_request:
                return JsonResponse(
                    {
                        "ok": True,
                        "created": created,
                        "name": saved_purpose.name,
                        "category": saved_purpose.category,
                        "purpose_count": len(category_purposes),
                        "purpose_options": category_purposes,
                        "message": (
                            "Expense purpose created successfully."
                            if created
                            else (
                                f"{saved_purpose.name} already exists for "
                                f"{saved_purpose.category}."
                            )
                        ),
                    }
                )

            if created:
                messages.success(request, "Expense purpose created successfully.")
            else:
                messages.info(
                    request,
                    (
                        f"{saved_purpose.name} already exists for "
                        f"{saved_purpose.category}."
                    ),
                )
            return redirect(self.get_next_url())

        if is_ajax_request:
            return JsonResponse(
                {
                    "ok": False,
                    "errors": {
                        field_name: [
                            error["message"]
                            for error in field_errors
                        ]
                        for field_name, field_errors in form.errors.get_json_data().items()
                    },
                },
                status=400,
            )

        first_error = next(
            iter(
                next(iter(form.errors.values()), ["Unable to save the purpose right now."])
            ),
            "Unable to save the purpose right now.",
        )
        messages.error(request, first_error)
        return redirect(self.get_next_url())


class DailySettlementView(ModulePermissionRequiredMixin, TemplateView):
    permission_field = "allow_daily_settlement"
    permission_denied_message = "You do not have access to Daily Settlement."
    template_name = "tracker/daily_settlement.html"

    def user_can_edit_opening_balance(self):
        return bool(self.request.user.is_superuser)

    def get_selected_entry_date(self, params, filter_values):
        return get_settlement_entry_date(params, filter_values["selected_date"])

    def get_loaded_settlement(self, settlement_date):
        return filter_queryset_by_role(
            DailyCashSettlement.objects.all(),
            self.request.user,
        ).filter(settlement_date=settlement_date).first()

    def get_settlement_queryset(self, filter_values):
        return filter_queryset_by_role(
            DailyCashSettlement.objects.all(),
            self.request.user,
        ).filter(settlement_date=filter_values["selected_date"]).order_by(
            "-settlement_date", "-updated_at", "-pk"
        )

    def get_credit_bill_queryset(self, settlement_date):
        return get_credit_bill_records_for_date(settlement_date)

    def build_form(self, selected_entry_date, loaded_settlement, autofill_summary):
        if loaded_settlement:
            preview_values = self.get_loaded_settlement_preview_values(
                loaded_settlement,
                autofill_summary,
            )
            preview = build_settlement_preview(
                preview_values,
                cash_denomination_total=loaded_settlement.cash_denomination_total,
            )
            form = DailyCashSettlementForm(
                instance=loaded_settlement,
                initial={
                    "opening_balance": preview_values["opening_balance"],
                    "gpay_settled": preview_values["gpay_settled"],
                    "closing_balance": preview["closing_balance"],
                },
            )
        else:
            preview = build_settlement_preview(
                {
                    "opening_balance": autofill_summary["opening_balance"],
                    "sales_ledger_cash": autofill_summary["sales_ledger_cash"],
                    "counter_income_amount": autofill_summary["counter_income_amount"],
                    "gpay_settled": autofill_summary["gpay_settled"],
                    "cash_settled": Decimal("0.00"),
                    "expense_amount": autofill_summary["expense_amount"],
                },
                cash_denomination_total=Decimal("0.00"),
            )
            form = DailyCashSettlementForm(
                initial={
                    "settlement_date": selected_entry_date,
                    "opening_balance": autofill_summary["opening_balance"],
                    "gpay_settled": autofill_summary["gpay_settled"],
                    "cash_settled": Decimal("0.00"),
                    "closing_balance": preview["closing_balance"],
                }
            )
        return self.configure_settlement_form(form, autofill_summary)

    def get_loaded_settlement_preview_values(self, loaded_settlement, autofill_summary):
        if autofill_summary["sales_count"] > 0:
            sales_ledger_cash = autofill_summary["sales_ledger_cash"]
            gpay_settled = autofill_summary["gpay_settled"]
        else:
            sales_ledger_cash = (
                loaded_settlement.actual_sales - loaded_settlement.gpay_settled
            )
            gpay_settled = autofill_summary["gpay_settled"]

        return {
            "opening_balance": loaded_settlement.opening_balance,
            "sales_ledger_cash": sales_ledger_cash,
            "counter_income_amount": autofill_summary["counter_income_amount"],
            "gpay_settled": gpay_settled,
            "cash_settled": loaded_settlement.cash_settled,
            "expense_amount": autofill_summary["expense_amount"],
        }

    def configure_settlement_form(self, form, autofill_summary):
        opening_class = form.fields["opening_balance"].widget.attrs.get("class", "")
        if self.user_can_edit_opening_balance():
            form.fields["opening_balance"].widget.attrs.pop("readonly", None)
            form.fields["opening_balance"].widget.attrs["class"] = opening_class.replace(
                "settlement-readonly", ""
            ).strip()
        else:
            form.fields["opening_balance"].widget.attrs["readonly"] = True
            form.fields["opening_balance"].widget.attrs["class"] = (
                f"{opening_class} settlement-readonly".strip()
            )

        closing_class = form.fields["closing_balance"].widget.attrs.get("class", "")
        form.fields["closing_balance"].widget.attrs["readonly"] = True
        form.fields["closing_balance"].widget.attrs["class"] = (
            f"{closing_class} settlement-readonly".strip()
        )

        current_class = form.fields["gpay_settled"].widget.attrs.get("class", "")
        form.fields["gpay_settled"].widget.attrs["readonly"] = True
        form.fields["gpay_settled"].widget.attrs["class"] = (
            f"{current_class} settlement-readonly".strip()
        )
        return form

    def get_cash_denominations(self, params=None, loaded_settlement=None):
        if params is not None and "cash_denominations" in params:
            return parse_cash_denominations_payload(params.get("cash_denominations"))
        if loaded_settlement:
            return parse_cash_denominations_payload(
                build_cash_denominations_payload(loaded_settlement.cash_denominations)
            )
        return {}

    def get_preview_source(self, form, loaded_settlement, autofill_summary):
        if form.is_bound:
            return {
                "opening_balance": (
                    form["opening_balance"].value()
                    if self.user_can_edit_opening_balance()
                    else (
                        loaded_settlement.opening_balance
                        if loaded_settlement
                        else autofill_summary["opening_balance"]
                    )
                ),
                "sales_ledger_cash": autofill_summary["sales_ledger_cash"],
                "counter_income_amount": autofill_summary["counter_income_amount"],
                "gpay_settled": form["gpay_settled"].value(),
                "cash_settled": form["cash_settled"].value(),
                "expense_amount": autofill_summary["expense_amount"],
            }
        if loaded_settlement:
            return self.get_loaded_settlement_preview_values(
                loaded_settlement,
                autofill_summary,
            )
        return {
            "opening_balance": autofill_summary["opening_balance"],
            "sales_ledger_cash": autofill_summary["sales_ledger_cash"],
            "counter_income_amount": autofill_summary["counter_income_amount"],
            "gpay_settled": autofill_summary["gpay_settled"],
            "cash_settled": Decimal("0.00"),
            "expense_amount": autofill_summary["expense_amount"],
        }

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        filter_values = kwargs.pop(
            "filter_values",
            get_settlement_filter_values(self.request.GET),
        )
        selected_entry_date = kwargs.pop(
            "selected_entry_date",
            self.get_selected_entry_date(self.request.GET, filter_values),
        )
        loaded_settlement = kwargs.pop(
            "loaded_settlement",
            self.get_loaded_settlement(selected_entry_date),
        )
        cash_denominations = kwargs.pop(
            "cash_denominations",
            self.get_cash_denominations(loaded_settlement=loaded_settlement),
        )
        cash_denomination_total = get_cash_denominations_total(cash_denominations)
        autofill_summary = build_settlement_autofill_summary(
            self.request.user,
            selected_entry_date,
        )
        form = kwargs.pop("form", None)
        if form is None:
            form = self.build_form(
                selected_entry_date,
                loaded_settlement,
                autofill_summary,
            )
        else:
            form = self.configure_settlement_form(form, autofill_summary)

        settlement_records = self.get_settlement_queryset(filter_values)
        credit_bill_records = self.get_credit_bill_queryset(selected_entry_date)
        settlement_preview = build_settlement_preview(
            self.get_preview_source(form, loaded_settlement, autofill_summary),
            cash_denomination_total=cash_denomination_total,
        )
        if loaded_settlement and not form.is_bound:
            autofill_summary = {
                **autofill_summary,
                "opening_balance": loaded_settlement.opening_balance,
                "gpay_settled": settlement_preview["gpay_settled"],
                "cash_in_hand": settlement_preview["cash_in_hand"],
            }
            cash_difference = get_cash_difference_amount(
                cash_denomination_total,
                settlement_preview["cash_in_hand"],
            )
        else:
            cash_difference = get_cash_difference_amount(
                cash_denomination_total,
                settlement_preview["cash_in_hand"],
            )
        settlement_totals = settlement_records.aggregate(
            gpay_total=Sum("gpay_settled"),
            cash_total=Sum("cash_settled"),
            expense_total=Sum("expense_amount"),
            actual_sales_total=Sum("actual_sales"),
            closing_total=Sum("closing_balance"),
        )
        credit_bill_total = sum(
            (record.effective_balance_amount for record in credit_bill_records),
            Decimal("0.00"),
        )

        context.update(
            {
                "form": form,
                "settlement_filters": filter_values,
                "selected_entry_date": selected_entry_date,
                "selected_settlement": loaded_settlement,
                "settlement_preview": settlement_preview,
                "can_edit_opening_balance": self.user_can_edit_opening_balance(),
                "settlement_records": settlement_records,
                "settlement_count": settlement_records.count(),
                "gpay_total": settlement_totals["gpay_total"] or Decimal("0.00"),
                "cash_total": settlement_totals["cash_total"] or Decimal("0.00"),
                "expense_total": settlement_totals["expense_total"] or Decimal("0.00"),
                "actual_sales_total": settlement_totals["actual_sales_total"]
                or Decimal("0.00"),
                "closing_total": settlement_totals["closing_total"] or Decimal("0.00"),
                "autofill_summary": autofill_summary,
                "cash_denomination_values": CASH_DENOMINATION_VALUES,
                "cash_denominations": cash_denominations,
                "cash_denomination_rows": build_cash_denomination_rows(
                    cash_denominations
                ),
                "cash_denominations_json": build_cash_denominations_payload(
                    cash_denominations
                ),
                "cash_denomination_total": cash_denomination_total,
                "cash_difference": cash_difference,
                "cash_difference_abs": abs(cash_difference),
                "credit_bill_records": credit_bill_records[:10],
                "credit_bill_count": len(credit_bill_records),
                "credit_bill_total": credit_bill_total,
            }
        )
        return context

    def post(self, request, *args, **kwargs):
        filter_values = get_settlement_filter_values(request.POST)
        selected_entry_date = self.get_selected_entry_date(request.POST, filter_values)
        loaded_settlement = self.get_loaded_settlement(selected_entry_date)
        cash_denominations = self.get_cash_denominations(
            params=request.POST,
            loaded_settlement=loaded_settlement,
        )
        cash_denomination_total = get_cash_denominations_total(cash_denominations)
        selected_date_autofill = build_settlement_autofill_summary(
            request.user,
            selected_entry_date,
        )
        post_data = request.POST.copy()
        if not request.user.is_superuser and not (post_data.get("opening_balance") or "").strip():
            post_data["opening_balance"] = str(
                loaded_settlement.opening_balance
                if loaded_settlement
                else selected_date_autofill["opening_balance"]
            )
        form = DailyCashSettlementForm(post_data, instance=loaded_settlement)

        if form.is_valid():
            settlement = form.save(commit=False)
            if settlement.pk is None:
                settlement.user = request.user
            autofill_summary = build_settlement_autofill_summary(
                request.user,
                settlement.settlement_date,
            )
            if request.user.is_superuser:
                settlement.opening_balance = form.cleaned_data["opening_balance"]
            elif loaded_settlement:
                settlement.opening_balance = loaded_settlement.opening_balance
            else:
                settlement.opening_balance = autofill_summary["opening_balance"]
            settlement.expense_amount = autofill_summary["expense_amount"]
            if autofill_summary["settlement_source"] == "sales":
                settlement.gpay_settled = autofill_summary["gpay_settled"]
            cash_in_hand = get_cash_in_hand_amount(
                settlement.opening_balance,
                autofill_summary["sales_ledger_cash"],
                settlement.expense_amount,
                counter_income_amount=autofill_summary["counter_income_amount"],
            )
            closing_balance = get_settlement_closing_balance(
                settlement.opening_balance,
                autofill_summary["sales_ledger_cash"],
                settlement.expense_amount,
                settlement.cash_settled,
                cash_denomination_total,
                counter_income_amount=autofill_summary["counter_income_amount"],
            )
            if closing_balance < 0:
                if cash_denominations:
                    error_message = (
                        "Cash settled cannot be greater than cash denomination total "
                        f"of Rs. {format_money(cash_denomination_total)}."
                    )
                else:
                    error_message = (
                        "Cash settled cannot be greater than cash in hand "
                        f"of Rs. {format_money(cash_in_hand)}."
                    )
                form.add_error(
                    "cash_settled",
                    error_message,
                )
                return self.render_to_response(
                    self.get_context_data(
                        form=form,
                        filter_values=filter_values,
                        selected_entry_date=selected_entry_date,
                        loaded_settlement=loaded_settlement,
                        cash_denominations=cash_denominations,
                    )
                )
            settlement.closing_balance = closing_balance
            settlement._expected_cash_in_hand = cash_in_hand
            settlement._expected_total_amount = (
                settlement.opening_balance
                + autofill_summary["sales_ledger_cash"]
                + settlement.gpay_settled
            )
            settlement._expected_actual_sales = (
                settlement._expected_total_amount - settlement.opening_balance
            )
            settlement.cash_denominations = cash_denominations
            settlement.save()
            messages.success(
                request,
                f"Daily cash settlement saved for {settlement.settlement_date:%d-%m-%Y}.",
            )
            try:
                recipients = send_daily_settlement_email(settlement)
            except Exception:
                logger.exception(
                    "Could not send daily settlement email for settlement %s.",
                    settlement.pk,
                )
                messages.warning(
                    request,
                    "Settlement was saved, but the email could not be sent.",
                )
            else:
                if recipients:
                    messages.info(
                        request,
                        "Daily settlement email sent to "
                        f"{', '.join(recipients)}.",
                    )
            return redirect(build_settlement_redirect_url(settlement.settlement_date))

        messages.error(request, "Please correct the highlighted settlement details.")
        return self.render_to_response(
            self.get_context_data(
                form=form,
                filter_values=filter_values,
                selected_entry_date=selected_entry_date,
                loaded_settlement=loaded_settlement,
                cash_denominations=cash_denominations,
            )
        )


class PurchaseListView(ModulePermissionRequiredMixin, AutoLoadPaginatedListView):
    permission_field = "allow_purchases"
    permission_denied_message = "You do not have access to Purchases."
    model = PurchaseRecord
    template_name = "tracker/purchase_list.html"
    context_object_name = "records"
    paginate_by = 20

    def get_queryset(self):
        queryset = get_purchase_base_queryset(self.request.user)
        queryset = apply_purchase_filters(
            queryset,
            get_purchase_filter_values(self.request.GET),
        )
        return queryset.order_by("-transaction_date", "-created_at", "-pk")

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip()
        if action != "sync_purchases":
            return redirect(build_purchase_redirect_url(request.POST))

        purchase_filters = get_purchase_filter_values(request.POST)
        try:
            stats = sync_purchases_from_sqlserver(
                user=request.user,
                date_from=parse_date(purchase_filters["date_from"]),
                date_to=parse_date(purchase_filters["date_to"]),
                supplier_name=purchase_filters["supplier_name"],
                invoice_number=purchase_filters["invoice_number"],
            )
        except PurchaseSyncError as exc:
            messages.error(request, str(exc))
        else:
            messages.success(
                request,
                "Purchase sync completed. "
                f"{stats.fetched_count} rows processed, "
                f"{stats.inserted_count} inserted, "
                f"{stats.refreshed_count} refreshed, "
                f"{stats.skipped_count} skipped.",
            )
        return redirect(build_purchase_redirect_url(purchase_filters))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        page_obj = context.get("page_obj")
        base_queryset = get_purchase_base_queryset(self.request.user)
        records = self.get_queryset()
        synced_purchase_queryset = filter_queryset_by_role(
            PurchaseRecord.objects.filter(
                source_reference__startswith=SQLSERVER_PURCHASE_SOURCE_PREFIX
            ),
            self.request.user,
        )
        raw_filter_values = get_raw_purchase_filter_values(self.request.GET)
        filter_values = get_purchase_filter_values(self.request.GET)
        today = date.today()
        month_queryset = base_queryset.filter(
            transaction_date__year=today.year,
            transaction_date__month=today.month,
        )
        today_queryset = base_queryset.filter(transaction_date=today)
        filtered_summary = summarize_purchase_queryset(records)
        context["purchase_filters"] = filter_values
        context["has_active_filters"] = any(raw_filter_values.values())
        context["is_default_today_view"] = not context["has_active_filters"]
        context["default_view_date"] = today
        context["month_summary"] = summarize_purchase_queryset(month_queryset)
        context["today_summary"] = summarize_purchase_queryset(today_queryset)
        context["filtered_summary"] = filtered_summary
        context["page_count"] = filtered_summary["count"]
        context["page_total"] = filtered_summary["total_amount"]
        context["paid_total"] = filtered_summary["paid_amount"]
        context["pending_total"] = filtered_summary["pending_amount"]
        context["synced_purchase_count"] = synced_purchase_queryset.count()
        context["recent_records"] = list(records[:8])
        context["previous_page_url"] = ""
        context["pagination_links"] = []

        if page_obj:
            if page_obj.has_previous():
                context["previous_page_url"] = build_page_url(
                    self.request,
                    page_obj.previous_page_number(),
                )
            if page_obj.has_next():
                context["next_page_url"] = build_page_url(
                    self.request,
                    page_obj.next_page_number(),
                )
            else:
                context["next_page_url"] = ""

            context["pagination_links"] = [
                {
                    "number": page_number,
                    "url": build_page_url(self.request, page_number),
                    "is_current": page_number == page_obj.number,
                }
                for page_number in page_obj.paginator.page_range
            ]
            context["showing_from"] = page_obj.start_index()
            context["showing_to"] = page_obj.end_index()
        else:
            context["showing_from"] = 1 if context["page_count"] else 0
            context["showing_to"] = context["page_count"]

        return context


class SalesListView(ModulePermissionRequiredMixin, AutoLoadPaginatedListView):
    permission_field = "allow_sales"
    permission_denied_message = "You do not have access to Sales."
    model = SalesLedgerRecord
    template_name = "tracker/sales_list.html"
    context_object_name = "records"

    def get_base_queryset(self):
        return SalesLedgerRecord.objects.filter(is_cancelled=False)

    def get_queryset(self):
        queryset = apply_sales_filters(
            self.get_base_queryset(),
            get_sales_filter_values(self.request.GET),
        )
        return queryset.order_by("-sale_date", "-source_sale_no")

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip()
        if action == "save_split":
            return self.handle_split_save(request)

        sales_filters = get_sales_filter_values(request.POST)
        try:
            stats = sync_sales_from_sqlserver(
                date_from=parse_date(sales_filters["date_from"]),
                date_to=parse_date(sales_filters["date_to"]),
            )
        except SalesSyncError as exc:
            messages.error(request, str(exc))
        else:
            messages.success(
                request,
                "Sales sync completed. "
                f"{stats.fetched_count} rows processed, "
                f"{stats.inserted_count} inserted, "
                f"{stats.refreshed_count} refreshed.",
            )
        return redirect(build_sales_redirect_url(sales_filters))

    def handle_split_save(self, request):
        record_id = (request.POST.get("record_id") or "").strip()
        record = self.get_base_queryset().filter(pk=record_id).first()
        if record is None:
            messages.error(request, "Sales bill not found.")
            return redirect(build_sales_redirect_url(request.POST))

        split_cash_amount = parse_money_value(request.POST.get("split_cash_amount"))
        split_card_amount = parse_money_value(request.POST.get("split_card_amount"))

        if split_cash_amount < 0 or split_card_amount < 0:
            messages.error(request, "Cash and card amounts cannot be negative.")
            return redirect(build_sales_redirect_url(request.POST))

        split_total = split_cash_amount + split_card_amount
        if split_total > record.net_amount:
            messages.error(
                request,
                (
                    f"Split total for bill {record.bill_no} cannot exceed "
                    f"the net amount of Rs. {format_money(record.net_amount)}."
                ),
            )
            return redirect(build_sales_redirect_url(request.POST))

        record.split_cash_amount = split_cash_amount
        record.split_card_amount = split_card_amount
        record.save(update_fields=["split_cash_amount", "split_card_amount"])

        if split_total == 0:
            messages.success(
                request,
                f"Split values cleared for bill {record.bill_no}.",
            )
        else:
            messages.success(
                request,
                (
                    f"Split saved for bill {record.bill_no}: "
                    f"Cash Rs. {format_money(split_cash_amount)} and "
                    f"Card Rs. {format_money(split_card_amount)}."
                ),
            )
        return redirect(build_sales_redirect_url(request.POST))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        base_queryset = self.get_base_queryset()
        filtered_queryset = self.get_queryset()
        raw_filter_values = get_raw_sales_filter_values(self.request.GET)
        filter_values = get_sales_filter_values(self.request.GET)
        today = date.today()
        visible_records = list(context.get("records") or [])
        filtered_records = list(filtered_queryset)
        filtered_summary = summarize_sales_records(filtered_records)
        credit_records = [
            record for record in filtered_records if get_effective_sales_balance_amount(record) > 0
        ]
        context["records"] = visible_records
        context["sales_filters"] = filter_values
        context["has_active_filters"] = any(raw_filter_values.values())
        context["is_default_today_view"] = not context["has_active_filters"]
        context["default_view_date"] = today
        context["month_summary"] = summarize_sales_records(
            base_queryset.filter(
                sale_date__year=today.year,
                sale_date__month=today.month,
            )
        )
        context["today_summary"] = summarize_sales_records(
            base_queryset.filter(sale_date=today)
        )
        context["filtered_summary"] = filtered_summary
        context["page_count"] = filtered_summary["count"]
        context["page_total"] = filtered_summary["total_amount"]
        context["received_total"] = filtered_summary["received_amount"]
        context["balance_total"] = filtered_summary["balance_amount"]
        context["split_cash_total"] = filtered_summary["split_cash_amount"]
        context["split_card_total"] = filtered_summary["split_card_amount"]
        context["credit_bill_records"] = credit_records[:10]
        context["credit_bill_count"] = filtered_summary["credit_count"]
        context["credit_bill_total"] = filtered_summary["credit_total"]
        context["payment_mode_options"] = [
            *SalesPaymentMode.choices[:-1],
            (SPLIT_PAYMENT_MODE_FILTER, SPLIT_PAYMENT_MODE_FILTER),
            SalesPaymentMode.choices[-1],
        ]
        context["synced_record_count"] = base_queryset.count()
        context["last_synced_at"] = (
            SalesLedgerRecord.objects.order_by("-synced_at")
            .values_list("synced_at", flat=True)
            .first()
        )
        return context


class PurchaseDetailView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_purchases"
    permission_denied_message = "You do not have access to Purchases."
    def get(self, request, pk, *args, **kwargs):
        purchase = get_object_or_404(
            get_purchase_base_queryset(request.user).prefetch_related("payments__user"),
            pk=pk,
        )
        return JsonResponse(build_purchase_detail_payload(purchase))


class PurchaseCreateView(ModulePermissionRequiredMixin, CreateView):
    permission_field = "allow_purchases"
    permission_denied_message = "You do not have access to Purchases."
    model = PurchaseRecord
    form_class = PurchaseForm
    template_name = "tracker/purchase_form.html"
    success_url = reverse_lazy("purchase-list")

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["user"] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.user = self.request.user
        self.object = form.save()
        if self.object.paid_amount > 0:
            create_purchase_payment(
                purchase=self.object,
                user=self.request.user,
                amount=self.object.paid_amount,
                notes="Initial paid amount saved with the purchase record.",
                payment_date=self.object.transaction_date,
                update_totals=False,
            )
        sync_purchase_to_expense(self.object)
        messages.success(self.request, "Purchase record saved successfully and added to expenses.")
        if "_download_invoice" in self.request.POST:
            return build_purchase_invoice_response(self.object)
        return redirect(self.get_success_url())

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["supplier_count"] = filter_queryset_by_role(
            Supplier.objects.all(),
            self.request.user,
        ).count()
        return context


class PurchaseInvoiceDownloadView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_purchases"
    permission_denied_message = "You do not have access to Purchases."
    def get(self, request, pk, *args, **kwargs):
        purchase = get_object_or_404(
            get_purchase_base_queryset(request.user),
            pk=pk,
        )
        return build_purchase_invoice_response(purchase)


class PurchasePaymentCreateView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_purchases"
    permission_denied_message = "You do not have access to Purchases."
    success_url = reverse_lazy("purchase-list")

    def post(self, request, pk, *args, **kwargs):
        with transaction.atomic():
            locked_queryset = filter_queryset_by_role(
                PurchaseRecord.objects.select_for_update(),
                request.user,
            )
            purchase = get_object_or_404(
                locked_queryset,
                pk=pk,
            )
            form = PurchasePaymentForm(request.POST, purchase=purchase)
            if not form.is_valid():
                error_text = " ".join(
                    error
                    for field_errors in form.errors.values()
                    for error in field_errors
                )
                messages.error(request, error_text or "Could not save the payment entry.")
                return redirect(self.success_url)

            create_purchase_payment(
                purchase=purchase,
                user=request.user,
                amount=form.cleaned_data["amount"],
                notes=form.cleaned_data["notes"],
            )
            sync_purchase_to_expense(purchase)

        messages.success(
            request,
            f"Payment of Rs. {format_money(form.cleaned_data['amount'])} saved for invoice {purchase.invoice_number}.",
        )
        return redirect(self.success_url)


class SupplierListView(ModulePermissionRequiredMixin, ListView):
    permission_field = "allow_suppliers"
    permission_denied_message = "You do not have access to Suppliers."
    model = Supplier
    template_name = "tracker/supplier_list.html"
    context_object_name = "suppliers"
    paginate_by = 30

    def get_base_queryset(self):
        return filter_queryset_by_role(Supplier.objects.all(), self.request.user)

    def get_search_query(self):
        return (self.request.GET.get("search") or "").strip()

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip()
        if action != "sync_suppliers":
            return redirect("supplier-list")

        try:
            stats = sync_suppliers_from_sqlserver(request.user)
        except SupplierSyncError as exc:
            messages.error(request, str(exc))
        else:
            messages.success(
                request,
                "Supplier sync completed. "
                f"{stats.fetched_count} rows processed, "
                f"{stats.inserted_count} inserted, "
                f"{stats.updated_count} updated, "
                f"{stats.skipped_count} skipped.",
            )
        return redirect(get_safe_next_url(request, reverse("supplier-list")))

    def get_queryset(self):
        queryset = self.get_base_queryset()
        search_query = self.get_search_query()
        if search_query:
            queryset = queryset.filter(
                Q(supplier_code__icontains=search_query)
                | Q(name__icontains=search_query)
                | Q(contact_person__icontains=search_query)
                | Q(phone_number__icontains=search_query)
                | Q(email__icontains=search_query)
            )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        suppliers = self.get_queryset()
        page_obj = context.get("page_obj")
        context["page_count"] = suppliers.count()
        context["active_supplier_count"] = suppliers.filter(
            status=SupplierStatus.ACTIVE
        ).count()
        context["inactive_supplier_count"] = suppliers.filter(
            status=SupplierStatus.INACTIVE
        ).count()
        context["synced_supplier_count"] = suppliers.exclude(
            source_supplier_no__isnull=True
        ).count()
        context["contact_person_count"] = suppliers.exclude(contact_person="").count()
        context["phone_number_count"] = suppliers.exclude(phone_number="").count()
        context["email_count"] = suppliers.exclude(email="").count()
        context["gstin_count"] = suppliers.exclude(gstin_number="").count()
        context["opening_balance_total"] = (
            suppliers.aggregate(total=Sum("opening_balance"))["total"]
            or Decimal("0.00")
        )
        context["search_query"] = self.get_search_query()
        context["current_url"] = self.request.get_full_path()
        context["previous_page_url"] = ""
        context["next_page_url"] = ""
        context["pagination_links"] = []
        if page_obj:
            if page_obj.has_previous():
                context["previous_page_url"] = build_page_url(
                    self.request, page_obj.previous_page_number()
                )
            if page_obj.has_next():
                context["next_page_url"] = build_page_url(
                    self.request, page_obj.next_page_number()
                )
            context["pagination_links"] = [
                {
                    "number": page_number,
                    "url": build_page_url(self.request, page_number),
                    "is_current": page_number == page_obj.number,
                }
                for page_number in page_obj.paginator.page_range
            ]
            context["showing_from"] = page_obj.start_index()
            context["showing_to"] = page_obj.end_index()
        else:
            context["showing_from"] = 1 if context["page_count"] else 0
            context["showing_to"] = context["page_count"]
        return context


class SupplierCreateView(ModulePermissionRequiredMixin, CreateView):
    permission_field = "allow_suppliers"
    permission_denied_message = "You do not have access to Suppliers."
    model = Supplier
    form_class = SupplierForm
    template_name = "tracker/supplier_form.html"
    success_url = reverse_lazy("supplier-list")

    def get_return_url(self):
        return get_safe_next_url(self.request, reverse("supplier-list"))

    def get_success_url(self):
        return self.get_return_url()

    def form_valid(self, form):
        form.instance.user = self.request.user
        messages.success(self.request, "Supplier details saved successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["next_supplier_code"] = get_next_supplier_code()
        context["next_url"] = self.get_return_url()
        context["is_editing"] = False
        return context


class SupplierUpdateView(ModulePermissionRequiredMixin, UpdateView):
    permission_field = "allow_suppliers"
    permission_denied_message = "You do not have access to Suppliers."
    model = Supplier
    form_class = SupplierForm
    template_name = "tracker/supplier_form.html"

    def get_queryset(self):
        return filter_queryset_by_role(Supplier.objects.all(), self.request.user)

    def get_return_url(self):
        return get_safe_next_url(self.request, reverse("supplier-list"))

    def get_success_url(self):
        return self.get_return_url()

    def form_valid(self, form):
        messages.success(self.request, "Supplier details updated successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["next_supplier_code"] = self.object.supplier_code or get_next_supplier_code()
        context["next_url"] = self.get_return_url()
        context["is_editing"] = True
        return context


class SupplierInactiveView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_suppliers"
    permission_denied_message = "You do not have access to Suppliers."

    def get_queryset(self):
        return filter_queryset_by_role(Supplier.objects.all(), self.request.user)

    def post(self, request, *args, **kwargs):
        supplier = get_object_or_404(self.get_queryset(), pk=kwargs["pk"])
        if supplier.status == SupplierStatus.INACTIVE:
            messages.info(request, "Supplier is already inactive.")
        else:
            supplier.status = SupplierStatus.INACTIVE
            supplier.save(update_fields=["status", "updated_at"])
            messages.success(request, "Supplier marked as inactive successfully.")
        return redirect(get_safe_next_url(request, reverse("supplier-list")))


class SupplierActiveView(ModulePermissionRequiredMixin, View):
    permission_field = "allow_suppliers"
    permission_denied_message = "You do not have access to Suppliers."

    def get_queryset(self):
        return filter_queryset_by_role(Supplier.objects.all(), self.request.user)

    def post(self, request, *args, **kwargs):
        supplier = get_object_or_404(self.get_queryset(), pk=kwargs["pk"])
        if supplier.status == SupplierStatus.ACTIVE:
            messages.info(request, "Supplier is already active.")
        else:
            supplier.status = SupplierStatus.ACTIVE
            supplier.save(update_fields=["status", "updated_at"])
            messages.success(request, "Supplier marked as active successfully.")
        return redirect(get_safe_next_url(request, reverse("supplier-list")))


class UserListView(AdminRequiredMixin, ListView):
    model = User
    template_name = "tracker/user_list.html"
    context_object_name = "users"
    paginate_by = 15

    def get_search_query(self):
        return (self.request.GET.get("search") or "").strip()

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip()
        if action and action != "sync_users":
            return redirect("user-list")
        try:
            stats = sync_users_from_sqlserver()
        except UserSyncError as exc:
            messages.error(request, str(exc))
        else:
            messages.success(
                request,
                "User sync completed. "
                f"{stats.fetched_count} rows processed, "
                f"{stats.inserted_count} inserted, "
                f"{stats.skipped_count} skipped.",
            )
        return redirect("user-list")

    def get_queryset(self):
        queryset = User.objects.select_related("account_profile", "module_permissions").order_by(
            "username"
        )
        search_query = self.get_search_query()
        if search_query:
            queryset = queryset.filter(
                Q(username__icontains=search_query)
                | Q(account_profile__master_name__icontains=search_query)
            )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        users = self.get_queryset()
        page_obj = context.get("page_obj")
        start_index = page_obj.start_index() if page_obj else 1
        for row_number, user in enumerate(context["users"], start=start_index):
            user.row_number = row_number
        for user in context["users"]:
            profile = getattr(user, "account_profile", None)
            user.master_name_display = (
                profile.master_name if profile and profile.master_name else "-"
            )
            user.role_label = get_user_role_label(user)
            user.status_label = "Active" if user.is_active else "Inactive"
            user.status_badge_class = (
                "user-status-badge-active" if user.is_active else "user-status-badge-inactive"
            )
            role_class_map = {
                "Admin": "user-role-badge-admin",
                "Store Admin": "user-role-badge-store-admin",
                "Staff": "user-role-badge-staff",
            }
            user.role_badge_class = role_class_map.get(
                user.role_label,
                "user-role-badge-staff",
            )

        context["page_count"] = users.count()
        context["active_count"] = users.filter(is_active=True).count()
        context["inactive_count"] = users.filter(is_active=False).count()
        context["admin_count"] = users.filter(is_superuser=True).count()
        context["store_admin_count"] = users.filter(
            is_superuser=False,
            is_staff=True,
        ).count()
        context["staff_count"] = users.filter(
            is_superuser=False,
            is_staff=False,
        ).count()
        context["master_name_count"] = users.filter(
            account_profile__isnull=False
        ).exclude(account_profile__master_name="").count()
        context["synced_user_count"] = users.filter(
            account_profile__source_user_no__isnull=False
        ).count()
        context["permission_count"] = users.filter(
            module_permissions__isnull=False
        ).count()
        context["search_query"] = self.get_search_query()
        context["current_url"] = self.request.get_full_path()
        context["previous_page_url"] = ""
        context["next_page_url"] = ""
        context["pagination_links"] = []
        if page_obj:
            if page_obj.has_previous():
                context["previous_page_url"] = build_page_url(
                    self.request, page_obj.previous_page_number()
                )
            if page_obj.has_next():
                context["next_page_url"] = build_page_url(
                    self.request, page_obj.next_page_number()
                )
            context["pagination_links"] = [
                {
                    "number": page_number,
                    "url": build_page_url(self.request, page_number),
                    "is_current": page_number == page_obj.number,
                }
                for page_number in page_obj.paginator.page_range
            ]
            context["showing_from"] = page_obj.start_index()
            context["showing_to"] = page_obj.end_index()
        else:
            context["showing_from"] = 1 if context["page_count"] else 0
            context["showing_to"] = context["page_count"]
        return context


class UserInactiveView(AdminRequiredMixin, View):
    def post(self, request, *args, **kwargs):
        managed_user = get_object_or_404(User.objects.all(), pk=kwargs["pk"])

        if managed_user.pk == request.user.pk:
            messages.error(request, "You cannot deactivate your own account.")
        elif not managed_user.is_active:
            messages.info(request, "User is already inactive.")
        else:
            managed_user.is_active = False
            managed_user.save(update_fields=["is_active"])
            messages.success(request, "User marked as inactive successfully.")

        return redirect(get_safe_next_url(request, reverse("user-list")))


class UserCreateView(AdminRequiredMixin, CreateView):
    model = User
    form_class = UserManagementForm
    template_name = "tracker/user_form.html"
    success_url = reverse_lazy("user-list")

    def get_return_url(self):
        return get_safe_next_url(self.request, reverse("user-list"))

    def get_success_url(self):
        return self.get_return_url()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["current_user"] = self.request.user
        kwargs["require_password"] = True
        return kwargs

    def form_valid(self, form):
        messages.success(self.request, "User created successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_eyebrow"] = "User Access"
        context["page_title"] = "Create User"
        context["page_description"] = (
            "Create admin, store admin, or staff users and control whether the account is active."
        )
        context["submit_label"] = "Save User"
        context["next_url"] = self.get_return_url()
        context["is_editing"] = False
        context["is_self_edit"] = False
        return context


class UserUpdateView(AdminRequiredMixin, UpdateView):
    model = User
    form_class = UserManagementForm
    template_name = "tracker/user_form.html"
    success_url = reverse_lazy("user-list")

    def get_return_url(self):
        return get_safe_next_url(self.request, reverse("user-list"))

    def get_success_url(self):
        return self.get_return_url()

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["current_user"] = self.request.user
        kwargs["require_password"] = False
        return kwargs

    def form_valid(self, form):
        messages.success(self.request, "User updated successfully.")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["page_eyebrow"] = "User Access"
        context["page_title"] = f"Edit User - {self.object.username}"
        context["page_description"] = (
            "Update role, active status, master name, and password for this account."
        )
        context["submit_label"] = "Update User"
        context["next_url"] = self.get_return_url()
        context["is_editing"] = True
        context["is_self_edit"] = self.object.pk == self.request.user.pk
        return context


class PermissionSettingsView(AdminRequiredMixin, TemplateView):
    template_name = "tracker/permission_settings.html"

    def get_selected_user(self):
        user_id = (self.request.GET.get("user") or "").strip()
        users = User.objects.filter(is_superuser=False).order_by("username")
        if user_id.isdigit():
            selected_user = users.filter(pk=user_id).first()
            if selected_user is not None:
                return selected_user
        return users.first()

    def post(self, request, *args, **kwargs):
        if request.headers.get("Content-Type", "").startswith("application/json"):
            try:
                payload = json.loads(request.body.decode("utf-8") or "{}")
            except (TypeError, ValueError):
                return JsonResponse(
                    {"ok": False, "message": "Invalid permission request."},
                    status=400,
                )
            user_id = str(payload.get("user", "")).strip()
            field_name = str(payload.get("field", "")).strip()
            field_names = {
                field for _label, field, _route in TRACKER_PERMISSION_ITEMS
            }

            if not user_id.isdigit():
                return JsonResponse(
                    {"ok": False, "message": "Select a valid user."},
                    status=400,
                )
            if field_name not in field_names:
                return JsonResponse(
                    {"ok": False, "message": "Invalid permission field."},
                    status=400,
                )

            managed_user = User.objects.filter(pk=user_id, is_superuser=False).first()
            if managed_user is None:
                return JsonResponse(
                    {"ok": False, "message": "User not found."},
                    status=404,
                )

            permission_record = get_user_permission_record(managed_user)
            raw_value = payload.get("value")
            is_enabled = raw_value if isinstance(raw_value, bool) else str(raw_value).strip().lower() in {
                "1",
                "true",
                "yes",
                "on",
            }
            setattr(permission_record, field_name, is_enabled)
            permission_record.save(update_fields=[field_name, "updated_at"])
            return JsonResponse({"ok": True, "value": is_enabled})

        messages.error(request, "Permission updates must be sent from the settings page.")
        return redirect("permission-settings")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        users = list(
            User.objects.filter(is_superuser=False)
            .select_related("account_profile", "module_permissions")
            .order_by("username")
        )
        selected_user = self.get_selected_user()
        permission_record = (
            get_user_permission_record(selected_user)
            if selected_user is not None
            else None
        )
        selected_profile = (
            getattr(selected_user, "account_profile", None)
            if selected_user is not None
            else None
        )

        permission_items = (
            build_permission_items(permission_record)
            if permission_record is not None
            else []
        )
        enabled_permission_count = sum(
            1 for item in permission_items if item["enabled"]
        )

        context.update(
            {
                "managed_users": users,
                "selected_user": selected_user,
                "selected_user_master_name": (
                    selected_profile.master_name
                    if selected_profile and selected_profile.master_name
                    else "-"
                ),
                "selected_user_role": (
                    get_user_role_label(selected_user) if selected_user is not None else ""
                ),
                "permission_items": permission_items,
                "enabled_permission_count": enabled_permission_count,
            }
        )
        return context


class ReconciliationWorkspaceView(ModulePermissionRequiredMixin, TemplateView):
    permission_field = "allow_reports"
    permission_denied_message = "You do not have access to Reconciliation."
    redirect_route_name = "reconciliation"
    include_history_records = True
    history_per_page = RECONCILIATION_HISTORY_PER_PAGE

    def get_redirect_url(self, params):
        return build_reconciliation_redirect_url(
            params,
            route_name=self.redirect_route_name,
        )

    def post(self, request, *args, **kwargs):
        action = (request.POST.get("action") or "").strip()
        if action == "save_opening_balance":
            filter_values = get_reconciliation_filter_values(request.POST)
            history_view = get_reconciliation_history_view(request.POST)
            opening_balance_form = ReconciliationOpeningBalanceForm(
                request.POST,
                prefix="opening",
            )
            if opening_balance_form.is_valid():
                save_reconciliation_opening_balance_for_date(
                    request.user,
                    opening_balance_form.cleaned_data["balance_date"],
                    opening_balance_form.cleaned_data["amount"],
                )
                messages.success(request, "Opening balance updated successfully.")
                return redirect(self.get_redirect_url(request.POST))
            messages.error(request, "Please correct the opening balance.")
            return self.render_to_response(
                self.get_context_data(
                    opening_balance_form=opening_balance_form,
                    filter_values=filter_values,
                    history_view=history_view,
                    opening_balance_modal_open=True,
                )
            )

        messages.info(
            request,
            "Add new records from the Income and Expenses pages. This screen only shows their history.",
        )
        return redirect(self.get_redirect_url(request.POST))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        filter_values = kwargs.pop(
            "filter_values",
            get_reconciliation_filter_values(self.request.GET),
        )
        history_view = kwargs.pop(
            "history_view",
            get_reconciliation_history_view(self.request.GET),
        )
        opening_balance_form = kwargs.pop(
            "opening_balance_form",
            ReconciliationOpeningBalanceForm(prefix="opening"),
        )
        opening_balance_modal_open = kwargs.pop("opening_balance_modal_open", False)

        reconciliation_opening_balance = get_reconciliation_opening_balance_for_date(
            self.request.user,
            filter_values["start_date"],
        )
        opening_balance_date = filter_values["start_date"]
        if opening_balance_form.is_bound:
            bound_balance_date = parse_date(
                (
                    opening_balance_form.data.get(
                        f"{opening_balance_form.prefix}-balance_date"
                    )
                    or ""
                ).strip()
            )
            if bound_balance_date is not None:
                opening_balance_date = bound_balance_date
        else:
            opening_balance_form.fields["balance_date"].initial = opening_balance_date
            opening_balance_form.initial["balance_date"] = opening_balance_date
            opening_balance_form.fields["amount"].initial = reconciliation_opening_balance
            opening_balance_form.initial["amount"] = reconciliation_opening_balance

        load_income_records = self.include_history_records and history_view == "income"
        load_expense_records = self.include_history_records and history_view == "expense"

        income_summary = build_reconciliation_income_entries(
            self.request.user,
            filter_values,
            page_number=(self.request.GET.get("income_page") or "").strip() or 1,
            per_page=self.history_per_page,
            include_records=load_income_records,
        )
        expense_summary = build_reconciliation_expense_entries(
            self.request.user,
            filter_values,
            page_number=(self.request.GET.get("expense_page") or "").strip() or 1,
            per_page=self.history_per_page,
            include_records=load_expense_records,
        )
        income_count = income_summary["record_count"]
        expense_count = expense_summary["record_count"]
        income_pagination = build_reconciliation_pagination(
            self.request,
            income_count,
            page_param="income_page",
            per_page=self.history_per_page,
            history_view="income",
        )
        expense_pagination = build_reconciliation_pagination(
            self.request,
            expense_count,
            page_param="expense_page",
            per_page=self.history_per_page,
            history_view="expense",
        )
        income_records = income_summary["records"] if load_income_records else []
        expense_records = expense_summary["records"] if load_expense_records else []
        manual_income_total = income_summary["manual_income_total"]
        settlement_income_total = income_summary["settlement_income_total"]
        split_card_balance_total = get_reconciliation_split_card_balance_for_date(
            self.request.user,
            filter_values["end_date"],
        )
        income_total = (
            reconciliation_opening_balance
            + manual_income_total
            + settlement_income_total
        )
        manual_expense_total = expense_summary["manual_expense_total"]
        manual_cash_expense_total = expense_summary["manual_cash_expense_total"]
        manual_non_cash_expense_total = expense_summary["manual_non_cash_expense_total"]
        purchase_cash_total = expense_summary["purchase_cash_total"]
        purchase_non_cash_total = expense_summary["purchase_non_cash_total"]
        purchase_total = expense_summary["purchase_total"]
        cash_expense_total = manual_cash_expense_total + purchase_cash_total
        non_cash_expense_total = (
            manual_non_cash_expense_total + purchase_non_cash_total
        )
        expense_total = manual_expense_total + purchase_total
        reconciliation_closing_balance = income_total - cash_expense_total
        next_day_opening_balance = reconciliation_closing_balance

        context.update(
            {
                "opening_balance_form": opening_balance_form,
                "opening_balance_modal_open": (
                    opening_balance_modal_open
                    or (opening_balance_form.is_bound and opening_balance_form.errors)
                ),
                "reconciliation_filters": {
                    "start_date": filter_values["start_date"].isoformat(),
                    "end_date": filter_values["end_date"].isoformat(),
                },
                "reconciliation_label": (
                    f"{filter_values['start_date']:%d-%m-%Y} "
                    f"to {filter_values['end_date']:%d-%m-%Y}"
                ),
                "reconciliation_history_view": history_view,
                "reconciliation_history_url": build_reconciliation_url(
                    filter_values,
                    history_view,
                    route_name="reconciliation",
                ),
                "reconciliation_summary_url": build_reconciliation_url(
                    filter_values,
                    history_view,
                    route_name="reconciliation-summary",
                ),
                "opening_balance_form_action_url": reverse(self.redirect_route_name),
                "income_history_loaded": load_income_records,
                "expense_history_loaded": load_expense_records,
                "income_records": income_records,
                "expense_records": expense_records,
                "income_page_obj": income_pagination["page_obj"],
                "expense_page_obj": expense_pagination["page_obj"],
                "income_previous_page_url": income_pagination["previous_page_url"],
                "expense_previous_page_url": expense_pagination["previous_page_url"],
                "income_next_page_url": income_pagination["next_page_url"],
                "expense_next_page_url": expense_pagination["next_page_url"],
                "income_pagination_links": income_pagination["pagination_links"],
                "expense_pagination_links": expense_pagination["pagination_links"],
                "income_showing_from": income_pagination["showing_from"],
                "income_showing_to": income_pagination["showing_to"],
                "expense_showing_from": expense_pagination["showing_from"],
                "expense_showing_to": expense_pagination["showing_to"],
                "reconciliation_opening_balance_date": opening_balance_date,
                "income_total": income_total,
                "settlement_income_total": settlement_income_total,
                "manual_income_total": manual_income_total,
                "manual_cash_income_total": income_summary["manual_cash_income_total"],
                "settlement_cash_total": income_summary["settlement_cash_total"],
                "settlement_card_total": income_summary["settlement_card_total"],
                "settlement_cash_target_display": income_summary[
                    "settlement_cash_target_display"
                ],
                "split_card_balance_total": split_card_balance_total,
                "reconciliation_opening_balance": reconciliation_opening_balance,
                "reconciliation_closing_balance": reconciliation_closing_balance,
                "next_day_opening_balance": next_day_opening_balance,
                "manual_expense_total": manual_expense_total,
                "manual_cash_expense_total": manual_cash_expense_total,
                "manual_non_cash_expense_total": manual_non_cash_expense_total,
                "purchase_cash_total": purchase_cash_total,
                "purchase_non_cash_total": purchase_non_cash_total,
                "purchase_total": purchase_total,
                "cash_expense_total": cash_expense_total,
                "non_cash_expense_total": non_cash_expense_total,
                "expense_total": expense_total,
                "net_total": income_total - expense_total,
                "income_count": income_count,
                "expense_count": expense_count,
            }
        )
        return context


class ReconciliationView(ReconciliationWorkspaceView):
    template_name = "tracker/reconciliation.html"


class ReconciliationSummaryView(ReconciliationWorkspaceView):
    template_name = "tracker/reconciliation_summary.html"
    redirect_route_name = "reconciliation-summary"
    include_history_records = False


class ReportsView(ModulePermissionRequiredMixin, TemplateView):
    permission_field = "allow_reports"
    permission_denied_message = "You do not have access to Reports."
    template_name = "tracker/reports.html"

    def get(self, request, *args, **kwargs):
        if (request.GET.get("export") or "").strip() == "excel":
            return build_reports_excel_response(request.user, request.GET)
        return super().get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(build_sales_report_workspace(self.request.user, self.request.GET))
        return context
